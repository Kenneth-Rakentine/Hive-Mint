"""The general spreadsheet builder: spec v4 of the detailer tracker rebuilds the legacy workbook exactly."""

import json
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from pypdf import PdfReader

from tools.builders import mobile_detailer, spreadsheet
from tools.spec import SpecError, prepare

V3 = Path(__file__).parent / "fixtures" / "detailer_spec_v3.json"


def _unquote(value):
    """'Jobs'!A1 and Jobs!A1 are the same reference."""
    return re.sub(r"'([A-Za-z_][A-Za-z0-9_]*)'!", lambda m: m.group(1) + "!", value) if isinstance(value, str) else value


def _cell(c):
    color = str(c.font.color.rgb) if c.font and c.font.color else None
    return (_unquote(c.value), c.number_format, bool(c.font and c.font.b), color,
            c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None, c.alignment.wrap_text, c.alignment.horizontal)


def _validations(ws):
    cells = set()
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            c1, r1, c2, r2 = range_boundaries(rng)
            cells |= {(dv.type, dv.operator, _unquote(dv.formula1), dv.formula2, c, r)
                      for c in range(c1, c2 + 1) for r in range(r1, r2 + 1)}
    return cells


@pytest.fixture(scope="module")
def both(tmp_path_factory, sample_spec):
    out = tmp_path_factory.mktemp("parity")
    legacy = mobile_detailer.build(prepare(json.loads(V3.read_text(encoding="utf-8"))), out / "legacy")
    generic = spreadsheet.build(sample_spec, out / "generic")
    return legacy, generic


def test_generic_workbook_matches_the_legacy_builder(both):
    (a, *_), (b, *_) = both
    wa, wb = load_workbook(a), load_workbook(b)
    assert wa.sheetnames == wb.sheetnames
    for name in wa.sheetnames:
        x, y = wa[name], wb[name]
        coords = {c.coordinate for ws in (x, y) for row in ws.iter_rows() for c in row if c.value is not None or c.has_style}
        assert [k for k in sorted(coords) if _cell(x[k]) != _cell(y[k])] == [], name
        assert _validations(x) == _validations(y), name
        assert {k: v.width for k, v in x.column_dimensions.items()} == {k: v.width for k, v in y.column_dimensions.items()}
        assert (x.freeze_panes, x.print_area, x.auto_filter.ref, x.sheet_view.showGridLines,
                x.page_setup.orientation) == (y.freeze_panes, y.print_area, y.auto_filter.ref,
                                              y.sheet_view.showGridLines, y.page_setup.orientation), name
        assert set(map(str, x.merged_cells.ranges)) == set(map(str, y.merged_cells.ranges))


def test_generic_charts_pdf_and_readme_match_the_legacy_builder(both):
    (wa, pa, ra), (wb, pb, rb) = both
    za, zb = zipfile.ZipFile(wa), zipfile.ZipFile(wb)
    charts = sorted(n for n in za.namelist() if "charts/chart" in n)
    assert len(charts) == 3 and all(za.read(n) == zb.read(n) for n in charts)
    text = [[p.extract_text() for p in PdfReader(path).pages] for path in (pa, pb)]
    assert text[0] == text[1]
    assert ra.read_bytes() == rb.read_bytes()


def test_unknown_tokens_and_names_are_spec_errors(tmp_path, sample_spec):
    spec = json.loads(json.dumps(sample_spec))
    jobs = next(s for s in spec["files"][0]["sheets"] if s["name"] == "Jobs")
    jobs["column_defs"][7]["formula"] = "=SUM({col:No Such Column}{r})"
    with pytest.raises(SpecError, match="No Such Column"):
        spreadsheet.build(spec, tmp_path / "a")
    jobs["column_defs"][7]["formula"] = "=SUM({range:Nowhere!Total})"
    with pytest.raises(SpecError, match="Nowhere"):
        spreadsheet.build(spec, tmp_path / "b")


def test_example_workbook_is_the_template_with_data_typed_in(tmp_path, sample_spec):
    spec = json.loads(json.dumps(sample_spec))
    base = spec["files"][0]["name"]
    spec["files"].insert(1, {"name": "Example.xlsx", "kind": "xlsx", "example_of": base, "role": "example",
                             "example_data": [{"sheet": "Jobs", "row": 1, "column": "Date", "value": "2026-03-05"},
                                              {"sheet": "Jobs", "row": 1, "column": "Service Amount", "value": 120},
                                              {"sheet": "Dashboard", "label": "Tax Year", "value": 2026}]})
    spec = prepare(spec)
    written = spreadsheet.build(spec, tmp_path / "out")
    assert [p.name for p in written] == [base, "Example.xlsx", "Instructions.pdf", "README.txt"]
    example = load_workbook(written[1])
    assert example["Jobs"]["F2"].value == 120 and example["Jobs"]["A2"].value.year == 2026
    assert example["Jobs"]["H2"].value.startswith("=IF(COUNT(")  # formulas kept
    assert len(example["Dashboard"]._charts) == 3
    template = load_workbook(written[0])
    assert template["Jobs"]["F2"].value is None
    assert next(f for f in spec["files"] if f["name"] == "Example.xlsx")["sheets"]  # normalize expanded it
