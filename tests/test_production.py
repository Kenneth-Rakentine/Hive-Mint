import json
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from agents.departments import STUB_HANDLERS
from agents.operations import OperationsManager
from agents.production import Production
from app.models import AgentRun, Build, Event
from app.states import Department, Outcome, ProductStatus as S
from tests.conftest import advance_to
from tools.packaging import sha256
from tools.spec import file_names


def _run_production(session, settings, product):
    ops = OperationsManager(session, settings, handlers={**STUB_HANDLERS, Department.PRODUCTION: Production(settings)})
    job = ops.next_job(product)
    ops.run_job(job)
    return job


def test_production_builds_every_spec_file_and_the_zip(session, settings, sample_product, sample_spec):
    advance_to(session, settings, sample_product, S.PRODUCTION)
    job = _run_production(session, settings, sample_product)

    assert sample_product.status == S.TESTING
    build = session.scalar(select(Build))
    out = settings.products_dir / str(sample_product.id) / "v1"
    assert build.version == 1 and build.job_id == job.id and build.path == str(out)
    assert sorted(p.name for p in out.iterdir()) == sorted(file_names(sample_spec) + [sample_spec["zip_name"]])
    with zipfile.ZipFile(out / sample_spec["zip_name"]) as archive:
        assert archive.namelist() == file_names(sample_spec)
    assert build.zip_sha256 == sha256(out / sample_spec["zip_name"])
    manifest = json.loads(build.manifest_json)
    assert [m["name"] for m in manifest] == file_names(sample_spec)
    assert all(m["sha256"] == sha256(out / m["name"]) for m in manifest)
    assert session.scalar(select(AgentRun).where(AgentRun.job_id == job.id)).provider == "deterministic"
    assert session.scalar(select(Event).where(Event.event_type == "build.created")) is not None
    assert not list(out.parent.glob(".*staging"))


def test_workbook_uses_live_formulas_and_buyer_friendly_layout(good_build, sample_spec):
    (file_spec,) = (f for f in sample_spec["files"] if f["kind"] == "xlsx")
    wb = load_workbook(good_build / file_spec["name"])
    assert wb.sheetnames[0] == "How to Use"
    assert [s["name"] for s in file_spec["sheets"]] == wb.sheetnames
    for sheet in (s for s in file_spec["sheets"] if "columns" in s):
        ws = wb[sheet["name"]]
        assert ws.freeze_panes == "A2"
        assert all(ws.column_dimensions[c.column_letter].width >= 8 for c in ws[1] if c.value)

    dashboard = wb["Dashboard"]
    cells = {c.value: dashboard.cell(c.row, c.column + 1) for row in dashboard.iter_rows() for c in row
             if isinstance(c.value, str)}
    assert str(cells["Net Profit"].value).startswith("=")
    assert cells["Tax Set-Aside %"].value == sample_spec["parameters"]["tax_set_aside_pct"]
    assert cells["Tax Set-Aside"].value == "=ROUND(MAX(0,B16)*B6,2)"
    assert wb["Jobs"]["H2"].value.startswith("=") and wb["Jobs"]["H2"].number_format == '"$"#,##0.00'
    assert wb["Jobs"]["A2"].number_format == "mm/dd/yyyy"
    assert len(dashboard._charts) == 3

    rates = wb["IRS Rates"]
    assert [rates.cell(r, 2).value for r in range(2, 2 + len(sample_spec["parameters"]["irs_mileage_rates"]))] == \
        [r["rate"] for r in sample_spec["parameters"]["irs_mileage_rates"]]


def test_one_job_entry_feeds_income_and_customer_history(good_build):
    """No double entry: customer rollups and Summary income both read the Jobs tab."""
    wb = load_workbook(good_build / "Mobile_Detailer_Business_Tracker.xlsx")
    customers = wb["Customers"]
    jobs_col = next(c.column for c in customers[1] if c.value == "Total Spent")
    assert "'Jobs'!" in customers.cell(2, jobs_col).value
    assert "'Jobs'!" in wb["Summary"]["B7"].value and "'Supplies'!" in wb["Summary"]["D7"].value


def test_instructions_pdf_and_readme_are_readable(good_build):
    reader = PdfReader(good_build / "Instructions.pdf")
    assert len(reader.pages) >= 2
    assert "Dashboard" in reader.pages[0].extract_text() + reader.pages[1].extract_text()
    assert "not tax advice" in (good_build / "README.txt").read_text(encoding="utf-8")


def test_rebuild_after_rework_creates_v2_and_keeps_v1(session, settings, sample_product):
    advance_to(session, settings, sample_product, S.PRODUCTION)
    _run_production(session, settings, sample_product)
    from app.lifecycle import transition

    transition(session, settings, sample_product, S.REWORK, actor="testing", reason="test failure")
    transition(session, settings, sample_product, S.PRODUCTION, actor="product_design", reason="rebuild")
    _run_production(session, settings, sample_product)

    versions = session.scalars(select(Build.version).order_by(Build.version)).all()
    assert versions == [1, 2]
    product_dir = settings.products_dir / str(sample_product.id)
    assert (product_dir / "v1").is_dir() and (product_dir / "v2").is_dir()


def test_unrecorded_folder_is_archived_not_overwritten(session, settings, sample_product):
    stale = settings.products_dir / str(sample_product.id) / "v1"
    stale.mkdir(parents=True)
    (stale / "old.txt").write_text("from a deleted database")
    advance_to(session, settings, sample_product, S.PRODUCTION)
    _run_production(session, settings, sample_product)

    assert sample_product.status == S.TESTING
    assert not (stale / "old.txt").exists()
    archived = list((settings.archive_dir / "stale-builds").glob("*/old.txt"))
    assert len(archived) == 1
    assert session.scalar(select(Event).where(Event.event_type == "build.stale_archived")) is not None


def test_product_without_spec_goes_to_rework(session, settings, make_product):
    product = advance_to(session, settings, make_product(), S.PRODUCTION)
    result = Production(settings)(product, None)
    assert result.outcome == Outcome.FAIL and result.next_status == S.REWORK
    assert "no specification" in result.summary
    assert not settings.products_dir.exists()
