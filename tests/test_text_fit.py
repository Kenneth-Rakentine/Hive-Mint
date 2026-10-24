"""QA catches boxed cells whose text runs past the box (the wedding Dashboard's couple name)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, PatternFill, Side

from agents.qa import check_text_fits
from tools.text_fit import overflows, text_width

BOX = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))


def _book(tmp_path, name="Book.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.column_dimensions["B"].width = 16
    ws["A1"] = "A long title that may spill into empty cells on purpose"            # unboxed: fine
    ws["B2"] = "='Start Here'!B5"                                                      # formula text, boxed
    ws["B2"].border = BOX
    ws["B3"] = "Short"
    ws["B3"].border = BOX
    ws["B4"] = "A very long boxed label here"
    ws["B4"].fill = PatternFill("solid", fgColor="EEF2F3")
    ws["B4"].alignment = Alignment(wrap_text=True)                                  # wraps: fine
    ws["B5"] = "Boxed and far too long text"
    ws["B5"].fill = PatternFill("solid", fgColor="EEF2F3")
    wb.save(tmp_path / name)
    return tmp_path / name


def test_text_width_grows_with_length_capitals_and_bold():
    assert text_width("Jordan & Sam (example)") > 16
    assert text_width("ABC") > text_width("abc") and text_width("abc", bold=True) > text_width("abc")


def test_only_boxed_unwrapped_text_wider_than_its_column_is_flagged(tmp_path):
    path = _book(tmp_path)
    found = overflows(path, {("DASHBOARD", "B2"): "Jordan & Sam (example)"})
    assert sorted(o.cell for o in found) == ["B2", "B5"]
    assert next(o for o in found if o.cell == "B2").width == 16


def test_qa_reports_overflow_as_a_major_finding(tmp_path):
    _book(tmp_path)
    spec = {"files": [{"name": "Book.xlsx", "kind": "xlsx"}]}
    findings = check_text_fits(spec, tmp_path, lambda p: {("DASHBOARD", "B2"): "Jordan & Sam (example)"})
    assert {f.where for f in findings} == {"Book.xlsx:Dashboard!B2", "Book.xlsx:Dashboard!B5"}
    assert all(f.severity == "major" and f.check == "text_fits" for f in findings)


def test_the_sample_product_has_no_overflowing_boxes(sample_spec, good_build):
    assert check_text_fits(sample_spec, good_build) == []
