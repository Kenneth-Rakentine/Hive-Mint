"""Testing department: the clean build passes every check, and each kind of breakage fails.

Expected files, sheets, columns and scenarios come from the product spec, not from here.
"""

import json
import zipfile

import pytest
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from sqlalchemy import select

from agents import testing
from agents.departments import STUB_HANDLERS
from agents.operations import OperationsManager
from agents.testing import Testing, check_files, check_pdf, check_text, check_workbook, check_zip
from app.models import Build, TestResult
from app.states import Department, JobStatus, ProductStatus as S
from config.settings import Settings
from tests.conftest import advance_to
from tools.packaging import make_zip, sha256
from tools.recalc import RecalcError, choose_engine
from tools.spec import file_names


def _file_spec(spec, name):
    return next(f for f in spec["files"] if f["name"] == name)


def _failed(checks):
    return {(c.check, c.target): c.detail for c in checks if not c.passed}


def _check_one_workbook(spec, build_dir, name):
    return check_workbook(spec, _file_spec(spec, name), spec["spreadsheet_rules"], build_dir / name,
                          engine="formulas", soffice=None)


def _edit(path, edit):
    wb = load_workbook(path)
    edit(wb)
    wb.save(path)


def _summary_cell(wb, label, offset=1):
    ws = wb["Summary"]
    cell = next(c for row in ws.iter_rows() for c in row if c.value == label)
    return ws.cell(cell.row, cell.column + offset)


# --------------------------------------------------------------------------- the clean build

def test_clean_build_passes_every_check(good_checks):
    assert _failed(good_checks) == {}


def test_every_spec_item_was_checked(good_checks, sample_spec):
    kinds = {c.check for c in good_checks}
    assert {"file_exists", "build_folder_clean", "zip_opens", "zip_contents", "zip_matches_build",
            "workbook_opens", "sheets_present", "columns_present", "header_frozen", "data_validation",
            "functions_compatible", "formula_text_clean", "no_error_values", "rate_table", "scenario",
            "pdf_opens", "pdf_page_count", "pdf_headings", "text_readable", "text_phrases"} <= kinds
    scenarios = [f"{f['name']}: {s['name']}" for f in sample_spec["files"] for s in f.get("scenarios", [])]
    assert sorted(c.target for c in good_checks if c.check == "scenario") == sorted(scenarios)
    assert {c.target for c in good_checks if c.check == "file_exists"} == set(file_names(sample_spec))


def test_spec_includes_the_milestone_scenarios(sample_spec):
    expectations = [(e.get("label") or e.get("column"), e["value"])
                    for f in sample_spec["files"] for s in f.get("scenarios", []) for e in s.get("expect", [])]
    assert ("Net Profit", 750) in expectations
    assert ("Deduction", 72.5) in expectations  # 100 miles x $0.725


# --------------------------------------------------------------------------- broken workbooks

def test_wrong_profit_formula_fails_the_profit_scenario(sample_spec, build_copy):
    def break_profit(wb):
        ws = wb["Summary"]
        for row in range(7, 19):
            ws[f"D{row}"] = f"=B{row}+C{row}"  # adds expenses instead of subtracting
    _edit(build_copy / "Income_Expense_Tracker.xlsx", break_profit)

    failed = _failed(_check_one_workbook(sample_spec, build_copy, "Income_Expense_Tracker.xlsx"))
    detail = next(d for (check, _), d in failed.items() if check == "scenario")
    assert "Net Profit" in detail and "expected 750" in detail and "1250" in detail


def test_division_by_zero_is_caught_on_the_empty_template(sample_spec, build_copy):
    _edit(build_copy / "Income_Expense_Tracker.xlsx",
          lambda wb: setattr(_summary_cell(wb, "Profit Margin"), "value", "=D19/B19"))  # no zero guard

    failed = _failed(_check_one_workbook(sample_spec, build_copy, "Income_Expense_Tracker.xlsx"))
    detail = failed[("no_error_values", "Income_Expense_Tracker.xlsx: empty template")]
    assert "#DIV/0!" in detail and "Summary!G10" in detail


def test_ref_value_and_name_errors_are_caught(sample_spec, build_copy):
    def break_cells(wb):
        ws = wb["Summary"]
        ws["J1"] = "=#REF!+1"
        ws["J2"] = '="abc"+1'
        ws["J3"] = "=NOSUCHFUNCTION(1)"
        ws["J4"] = "=XLOOKUP(1,A1:A2,B1:B2)"  # real in Excel 365, missing from older Excel
    _edit(build_copy / "Income_Expense_Tracker.xlsx", break_cells)

    failed = _failed(_check_one_workbook(sample_spec, build_copy, "Income_Expense_Tracker.xlsx"))
    template_errors = failed[("no_error_values", "Income_Expense_Tracker.xlsx: empty template")]
    for error in ("#REF!", "#VALUE!", "#NAME?"):
        assert error in template_errors
    assert "XLOOKUP" in failed[("functions_compatible", "Income_Expense_Tracker.xlsx")]
    assert "Summary!J1" in failed[("formula_text_clean", "Income_Expense_Tracker.xlsx")]


def test_missing_column_and_sheet_are_caught(sample_spec, build_copy):
    def break_layout(wb):
        wb["Income"]["G1"] = "Gratuity"  # spec requires "Tip"
        del wb["Lists"]
    _edit(build_copy / "Income_Expense_Tracker.xlsx", break_layout)

    failed = _failed(_check_one_workbook(sample_spec, build_copy, "Income_Expense_Tracker.xlsx"))
    assert "Tip" in failed[("columns_present", "Income_Expense_Tracker.xlsx:Income")]
    assert "Lists" in failed[("sheets_present", "Income_Expense_Tracker.xlsx")]


def test_unfrozen_header_and_missing_validation_are_caught(sample_spec, build_copy):
    def loosen(wb):
        wb["Expenses"].freeze_panes = None
        wb["Expenses"].data_validations.dataValidation = []
    _edit(build_copy / "Income_Expense_Tracker.xlsx", loosen)

    checks = check_workbook(sample_spec, _file_spec(sample_spec, "Income_Expense_Tracker.xlsx"),
                            sample_spec["spreadsheet_rules"], build_copy / "Income_Expense_Tracker.xlsx",
                            engine="formulas")
    failed = _failed(checks)
    assert ("header_frozen", "Income_Expense_Tracker.xlsx:Expenses") in failed
    assert "Category" in failed[("data_validation", "Income_Expense_Tracker.xlsx:Expenses")]


def test_wrong_irs_rate_is_caught(sample_spec, build_copy):
    _edit(build_copy / "Mileage_Log.xlsx", lambda wb: wb["IRS Rates"].cell(3, 2, 0.65))

    failed = _failed(_check_one_workbook(sample_spec, build_copy, "Mileage_Log.xlsx"))
    assert ("rate_table", "Mileage_Log.xlsx:IRS Rates") in failed
    assert any(check == "scenario" and "expected 72.5" in d for (check, _), d in failed.items())


def test_corrupt_workbook_fails_to_open(sample_spec, build_copy):
    (build_copy / "Customer_Log.xlsx").write_bytes(b"not a spreadsheet")
    failed = _failed(_check_one_workbook(sample_spec, build_copy, "Customer_Log.xlsx"))
    assert list(failed) == [("workbook_opens", "Customer_Log.xlsx")]


# --------------------------------------------------------------------------- files, ZIP, PDF, README

def test_missing_and_extra_files_are_caught(sample_spec, build_copy):
    (build_copy / "README.txt").unlink()
    (build_copy / "Thumbs.db").write_bytes(b"junk")
    failed = _failed(check_files(sample_spec, build_copy))
    assert ("file_exists", "README.txt") in failed
    assert "Thumbs.db" in failed[("build_folder_clean", "v1")]


@pytest.mark.parametrize("junk", ["__MACOSX/._Mileage_Log.xlsx", ".DS_Store", "build.log"])
def test_zip_with_build_junk_fails(sample_spec, build_copy, junk):
    zip_path = build_copy / sample_spec["zip_name"]
    recorded = sha256(zip_path)
    with zipfile.ZipFile(zip_path, "a") as archive:
        archive.writestr(junk, b"junk")
    failed = _failed(check_zip(sample_spec, build_copy, recorded))
    assert junk in failed[("zip_contents", sample_spec["zip_name"])]
    assert ("zip_integrity", sample_spec["zip_name"]) in failed


def test_zip_missing_a_file_or_out_of_date_fails(sample_spec, build_copy):
    files = [build_copy / n for n in file_names(sample_spec)]
    make_zip(build_copy / sample_spec["zip_name"], files[1:])
    failed = _failed(check_zip(sample_spec, build_copy, None))
    assert files[0].name in failed[("zip_contents", sample_spec["zip_name"])]

    make_zip(build_copy / sample_spec["zip_name"], files)
    (build_copy / "Instructions.pdf").write_bytes(b"%PDF changed")
    assert "Instructions.pdf" in _failed(check_zip(sample_spec, build_copy, None))[
        ("zip_matches_build", sample_spec["zip_name"])]


def test_not_a_zip_fails(sample_spec, build_copy):
    (build_copy / sample_spec["zip_name"]).write_bytes(b"not a zip")
    assert ("zip_opens", sample_spec["zip_name"]) in _failed(check_zip(sample_spec, build_copy, None))


def test_short_pdf_without_headings_fails(sample_spec, tmp_path):
    path = tmp_path / "Instructions.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Getting Started")
    pdf.save()
    failed = _failed(check_pdf(_file_spec(sample_spec, "Instructions.pdf"), path))
    assert "1 pages" in failed[("pdf_page_count", "Instructions.pdf")]
    missing = failed[("pdf_headings", "Instructions.pdf")]
    assert "Mileage Log" in missing and "Getting Started" not in missing


def test_unreadable_pdf_and_readme_missing_phrases_fail(sample_spec, tmp_path):
    (tmp_path / "Instructions.pdf").write_bytes(b"garbage")
    assert ("pdf_opens", "Instructions.pdf") in _failed(
        check_pdf(_file_spec(sample_spec, "Instructions.pdf"), tmp_path / "Instructions.pdf"))
    (tmp_path / "README.txt").write_text("Thanks for buying!")
    assert "not tax advice" in _failed(check_text(_file_spec(sample_spec, "README.txt"), tmp_path / "README.txt"))[
        ("text_phrases", "README.txt")]


# --------------------------------------------------------------------------- recalculation engine

def test_engine_falls_back_to_formulas_without_libreoffice(tmp_path):
    assert choose_engine("formulas", None) == ("formulas", None)
    assert choose_engine("auto", str(tmp_path / "missing.exe")) == ("formulas", None)
    with pytest.raises(RecalcError):
        choose_engine("libreoffice", str(tmp_path / "missing.exe"))
    fake = tmp_path / "soffice.exe"
    fake.write_bytes(b"")
    assert choose_engine("auto", str(fake)) == ("libreoffice", str(fake))


def test_unknown_engine_setting_is_rejected():
    with pytest.raises(ValueError):
        Settings(recalc_engine="excel")


# --------------------------------------------------------------------------- department + database

def _record_build(session, product, build_dir, spec):
    from app.models import Job

    job = Job(product_id=product.id, department=Department.PRODUCTION, job_type="build_product",
              for_status=S.PRODUCTION, status=JobStatus.SUCCEEDED)
    session.add(job)
    session.flush()
    build = Build(product_id=product.id, job_id=job.id, version=1, spec_slug=product.slug, path=str(build_dir),
                  zip_path=str(build_dir / spec["zip_name"]), zip_sha256=sha256(build_dir / spec["zip_name"]))
    session.add(build)
    session.flush()
    return build


def _testing_ops(session, settings):
    return OperationsManager(session, settings, handlers={**STUB_HANDLERS, Department.TESTING: Testing(settings)})


def test_broken_workbook_sends_product_to_rework_with_reasons(session, settings, sample_product, sample_spec,
                                                              build_copy):
    def break_profit(wb):
        for row in range(7, 19):
            wb["Summary"][f"D{row}"] = f"=B{row}+C{row}"
    _edit(build_copy / "Income_Expense_Tracker.xlsx", break_profit)
    make_zip(build_copy / sample_spec["zip_name"], [build_copy / n for n in file_names(sample_spec)])
    advance_to(session, settings, sample_product, S.TESTING)
    build = _record_build(session, sample_product, build_copy, sample_spec)

    ops = _testing_ops(session, settings)
    job = ops.next_job(sample_product)
    ops.run_job(job)

    assert sample_product.status == S.REWORK and sample_product.rework_count == 1
    assert "Testing failed" in sample_product.status_reason and "Net Profit" in sample_product.status_reason
    results = session.scalars(select(TestResult).where(TestResult.job_id == job.id)).all()
    assert results and all(r.build_id == build.id and r.product_id == sample_product.id for r in results)
    assert all(r.agent_run_id is not None for r in results)
    failed = [r for r in results if not r.passed]
    assert [r.check for r in failed] == ["scenario"]
    output = json.loads(job.output_json)
    assert output["outcome"] == "FAIL" and output["failed"] == 1 and output["engine"] == "formulas"


def test_passing_build_moves_product_to_qa(session, settings, sample_product, sample_spec, good_build, good_checks,
                                           monkeypatch):
    monkeypatch.setattr(testing, "run_checks", lambda *args, **kwargs: good_checks)  # already run once
    advance_to(session, settings, sample_product, S.TESTING)
    build = _record_build(session, sample_product, good_build, sample_spec)

    ops = _testing_ops(session, settings)
    job = ops.next_job(sample_product)
    ops.run_job(job)

    assert sample_product.status == S.QA
    results = session.scalars(select(TestResult).where(TestResult.build_id == build.id)).all()
    assert len(results) == len(good_checks) and all(r.passed for r in results)


def test_testing_without_a_build_fails_to_rework(session, settings, sample_product):
    advance_to(session, settings, sample_product, S.TESTING)
    ops = _testing_ops(session, settings)
    ops.run_job(ops.next_job(sample_product))
    assert sample_product.status == S.REWORK and "no build" in sample_product.status_reason
