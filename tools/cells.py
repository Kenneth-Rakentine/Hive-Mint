"""Locating cells the way specs address them, shared by the builder (example data), Testing
(scenario inputs and expectations) and QA.

A reference is either {sheet, row, column} (row 1 = first data row under the header) or
{sheet, label, offset?, row_offset?} (the cell `offset` columns right of a cell whose text is
`label`, default 1, shifted `row_offset` rows down).
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from tools.spec import sheet_spec


def headers(ws, header_row: int) -> dict[str, str]:
    """Header text -> column letter."""
    return {str(c.value).strip(): c.column_letter for c in ws[header_row] if c.value is not None}


def locate(wb: Workbook, file_spec: dict, ref: dict) -> str:
    """Spec reference -> cell coordinate. Raises LookupError with a readable reason."""
    if ref["sheet"] not in wb.sheetnames:
        raise LookupError(f"sheet '{ref['sheet']}' not found")
    ws = wb[ref["sheet"]]
    if "label" in ref:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() == ref["label"]:
                    col = cell.column + ref.get("offset", 1)
                    return f"{get_column_letter(col)}{cell.row + ref.get('row_offset', 0)}"
        raise LookupError(f"label '{ref['label']}' not found on {ref['sheet']}")
    sheet = sheet_spec(file_spec, ref["sheet"])
    header_row = sheet.get("header_row", 1)
    found = headers(ws, header_row)
    if ref["column"] not in found:
        raise LookupError(f"column '{ref['column']}' not found on {ref['sheet']}")
    return f"{found[ref['column']]}{header_row + ref['row']}"


def describe(ref: dict) -> str:
    if "label" in ref:
        offsets = "".join(f" {k}={ref[k]}" for k in ("row_offset", "offset") if k in ref)
        return f"{ref['sheet']} '{ref['label']}'{offsets}"
    return f"{ref['sheet']} row {ref['row']} '{ref['column']}'"
