"""A compact, text-only digest of a build for AI reviewers (local QA model, Claude Code Red Team).

Reviewers never get raw binary: only sheet names, table headers, instruction and label text,
chart titles, and the PDF and README text, trimmed to a character budget.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

DEFAULT_MAX_CHARS = 18_000
MAX_SHEET_TEXT = 3_000
MAX_PDF_TEXT = 9_000


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def workbook_text(file_spec: dict, path: Path) -> dict[str, dict]:
    """Per sheet: headers (table sheets), text cells (other sheets) and chart titles."""
    wb = load_workbook(path)
    tables = {s["name"] for s in file_spec.get("sheets", []) if "columns" in s}
    sheets = {}
    for ws in wb.worksheets:
        entry: dict = {"headers": [], "text": [], "charts": []}
        if ws.title in tables:
            entry["headers"] = [str(c.value).strip() for c in ws[1] if c.value is not None]
        else:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and not cell.value.startswith("="):
                        entry["text"].append(_clean(cell.value))
        for chart in ws._charts:
            try:
                entry["charts"].append("".join(r.t for p in chart.title.tx.rich.p for r in (p.r or [])))
            except AttributeError:
                entry["charts"].append("(untitled chart)")
        sheets[ws.title] = entry
    return sheets


def pdf_text(path: Path) -> str:
    reader = PdfReader(path)
    return _clean(" ".join(page.extract_text() or "" for page in reader.pages))


def product_digest(spec: dict, build_dir: Path, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    parts = [f"PRODUCT: {spec.get('name', spec['slug'])} (spec v{spec.get('spec_version', 1)})",
             f"FILES IN THE DOWNLOAD ({spec['zip_name']}): " + ", ".join(f["name"] for f in spec["files"])]
    if spec.get("features"):
        parts.append("PROMISED FEATURES: " + "; ".join(f["description"] for f in spec["features"]))
    for file_spec in spec["files"]:
        path = build_dir / file_spec["name"]
        if not path.is_file():
            parts.append(f"\n## {file_spec['name']}: MISSING")
            continue
        parts.append(f"\n## {file_spec['name']}")
        try:
            if file_spec["kind"] == "xlsx":
                for name, entry in workbook_text(file_spec, path).items():
                    lines = [f"### Sheet: {name}"]
                    if entry["headers"]:
                        lines.append("Columns: " + " | ".join(entry["headers"]))
                    if entry["text"]:
                        lines.append("Text: " + " / ".join(entry["text"])[:MAX_SHEET_TEXT])
                    if entry["charts"]:
                        lines.append("Charts: " + "; ".join(entry["charts"]))
                    parts.append("\n".join(lines))
            elif file_spec["kind"] == "pdf":
                text = pdf_text(path)
                parts.append(text[:MAX_PDF_TEXT] + (" [... trimmed]" if len(text) > MAX_PDF_TEXT else ""))
            else:
                parts.append(_clean(path.read_text(encoding="utf-8")))
        except Exception as error:  # an unreadable file is itself worth telling the reviewer
            parts.append(f"(could not read: {type(error).__name__}: {error})")
    digest = "\n".join(parts)
    if len(digest) > max_chars:
        digest = digest[:max_chars] + f"\n[... digest trimmed to {max_chars} characters]"
    return digest
