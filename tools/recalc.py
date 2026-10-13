"""Deterministic formula recalculation for Testing.

Uses LibreOffice headless when it is installed (closest to a real spreadsheet app),
otherwise the pure-Python `formulas` package. Either way the result is a plain dict of
computed values keyed by (SHEET NAME upper-cased, "A1").
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from tools.spec import excel_serial

Values = dict[tuple[str, str], object]

WINDOWS_SOFFICE = (
    Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "LibreOffice" / "program" / "soffice.exe",
    Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "LibreOffice" / "program" / "soffice.exe",
)
_SOLUTION_KEY = re.compile(r"^'\[[^\]]+\](?P<sheet>.+)'!(?P<cell>[A-Z]+[0-9]+)$")
SOFFICE_TIMEOUT_S = 120


class RecalcError(RuntimeError):
    """The workbook could not be recalculated."""


def find_soffice(configured: str | None = None) -> str | None:
    if configured:
        return configured if Path(configured).is_file() else None
    for name in ("soffice", "soffice.exe", "libreoffice"):
        if found := shutil.which(name):
            return found
    return next((str(p) for p in WINDOWS_SOFFICE if p.is_file()), None)


def choose_engine(preference: str, soffice_path: str | None) -> tuple[str, str | None]:
    """Resolve "auto" to a concrete engine. Returns (engine, soffice executable or None)."""
    soffice = find_soffice(soffice_path)
    if preference == "libreoffice" and soffice is None:
        raise RecalcError("HIVEMINT_RECALC_ENGINE=libreoffice but soffice was not found")
    if preference in ("auto", "libreoffice") and soffice:
        return "libreoffice", soffice
    return "formulas", None


def recalculate(path: Path, engine: str, workdir: Path, soffice: str | None = None) -> Values:
    if engine == "libreoffice":
        return _libreoffice(path, workdir, soffice)
    if engine == "formulas":
        return _formulas(path)
    raise RecalcError(f"Unknown recalculation engine: {engine}")


def _plain(value) -> object:
    """Normalize engine values: numpy scalars to Python, errors to their '#...' text, empties to None."""
    import schedula
    from formulas.tokens.operand import XlError

    if value is schedula.EMPTY:
        return None
    if isinstance(value, XlError):
        return str(value)
    if hasattr(value, "item"):  # numpy scalar
        value = value.item()
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return excel_serial(value)
    return value


def _formulas(path: Path) -> Values:
    os.environ.setdefault("TQDM_DISABLE", "1")  # formulas shows a progress bar while loading
    import formulas

    try:
        model = formulas.ExcelModel().loads(str(path)).finish()
        solution = model.calculate()
    except Exception as error:  # an engine crash is a test failure, not a pass
        raise RecalcError(f"formulas could not calculate {path.name}: {type(error).__name__}: {error}") from error
    values: Values = {}
    for key, ranges in solution.items():
        match = _SOLUTION_KEY.match(key)
        if not match:
            continue  # multi-cell range nodes
        raw = getattr(ranges, "value", ranges)
        try:
            raw = raw[0][0]
        except (TypeError, IndexError, KeyError):
            pass
        values[(match["sheet"].upper(), match["cell"])] = _plain(raw)
    return values


def _libreoffice(path: Path, workdir: Path, soffice: str | None) -> Values:
    if soffice is None:
        raise RecalcError("soffice not found")
    out_dir = workdir / "lo-out"
    out_dir.mkdir(parents=True, exist_ok=True)
    profile = (workdir / "lo-profile").resolve().as_uri()  # private profile: no clash with a running LibreOffice
    command = [soffice, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--calc",
               "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)]
    try:
        subprocess.run(command, check=True, capture_output=True, timeout=SOFFICE_TIMEOUT_S)
    except (subprocess.SubprocessError, OSError) as error:
        raise RecalcError(f"LibreOffice failed on {path.name}: {error}") from error
    converted = out_dir / path.name
    if not converted.is_file():
        raise RecalcError(f"LibreOffice produced no output for {path.name}")
    formulas_wb = load_workbook(path)
    values_wb = load_workbook(converted, data_only=True)
    values: Values = {}
    formula_cells = computed = 0
    for ws in values_wb.worksheets:
        source = formulas_wb[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, datetime):
                    value = value.date()
                if isinstance(value, date):
                    value = excel_serial(value)
                values[(ws.title.upper(), cell.coordinate)] = value
                if isinstance(source[cell.coordinate].value, str) and source[cell.coordinate].value.startswith("="):
                    formula_cells += 1
                    computed += value is not None
    if formula_cells and not computed:
        raise RecalcError(f"LibreOffice did not recalculate {path.name} (no cached formula results)")
    return values
