"""Deterministic chart checks (no AI): what a chart will look like, decided from its XML.

Each check catches a defect that once shipped (see rework brief 002), or one the spec rules out:
- chart_title_overlay: a title or legend drawn over the plot (overlay on, or missing, which
  Excel treats as on)
- chart_axis_title: an axis title the spec requires is missing
- chart_month_labels: month-axis labels longer than 3 characters, or without an explicit
  0 rotation (Excel then rotates them when space is short)
- chart_legend: a legend at the bottom, where it collides with category labels
- chart_number_format: a value axis without its own number format
- chart_axis_scale: a value axis without a fixed major unit, or without a minimum unless the
  spec allows negatives (an empty workbook then shows $0 $0 $1 $1); values must match the spec
- chart_overlap, chart_print_area, chart_width: anchors overlapping each other, outside the
  print area, or past the screen width the spec allows. Geometry is measured at each DPI in
  CHECK_DPIS, since Excel's column widths depend on the display.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from tools.chart_xml import CHECK_DPIS, ChartInfo, parse_ref, read_charts

MONTH_NAMES = {m.lower() for m in (
    "January February March April May June July August September October November December "
    "Jan Feb Mar Apr May Jun Jul Aug Sep Sept Oct Nov Dec").split()}
MAX_MONTH_LABEL = 3


@dataclass(frozen=True)
class ChartProblem:
    check: str
    where: str
    issue: str
    suggestion: str


def _where(file_name: str, chart: ChartInfo) -> str:
    return f"{file_name}:{chart.sheet}: chart {chart.index} '{chart.title or 'untitled'}'"


def category_labels(chart: ChartInfo, workbook) -> list[str]:
    """The category labels: the chart's cached labels, else the referenced cells' values."""
    if chart.category_cache:
        return chart.category_cache
    labels: list[str] = []
    for ref in chart.category_refs[:1]:
        parsed = parse_ref(ref)
        if parsed is None or parsed[0] not in workbook.sheetnames:
            continue
        sheet, cells = parsed
        for row in workbook[sheet][cells]:
            for cell in (row if isinstance(row, tuple) else (row,)):
                labels.append("" if cell.value is None else str(cell.value))
    return labels


def _is_month_axis(chart_spec: dict | None, labels: list[str]) -> bool:
    if chart_spec and chart_spec.get("category_axis") == "months":
        return True
    named = [label for label in labels if label]
    return bool(named) and all(re.sub(r"\W", "", label).lower() in MONTH_NAMES for label in named)


def _check_chart(chart: ChartInfo, chart_spec: dict | None, file_name: str, workbook) -> list[ChartProblem]:
    where = _where(file_name, chart)
    problems: list[ChartProblem] = []

    def add(check: str, issue: str, suggestion: str) -> None:
        problems.append(ChartProblem(check, where, issue, suggestion))

    if chart.has_title and chart.title_overlay is not False:
        add("chart_title_overlay", "The title is drawn over the plot area (overlay is "
            f"{'on' if chart.title_overlay else 'not set, which Excel treats as on'})",
            "Set the title overlay off (<c:overlay val=\"0\"/>) so the plot starts below the title")
    if chart.legend_pos is not None and chart.legend_overlay is not False:
        add("chart_title_overlay", "The legend is drawn over the plot area (overlay is "
            f"{'on' if chart.legend_overlay else 'not set, which Excel treats as on'})",
            "Set the legend overlay off (<c:overlay val=\"0\"/>)")

    category_axes = [a for a in chart.axes if a.kind in ("cat", "date") and not a.deleted]
    value_axes = [a for a in chart.axes if a.kind == "val" and not a.deleted]
    for kind in (chart_spec or {}).get("axis_titles", []):
        axes = value_axes if kind == "value" else category_axes
        if not any(a.has_title for a in axes):
            add("chart_axis_title", f"The {kind} axis has no title, but the spec requires one",
                f"Add a {kind} axis title")

    labels = category_labels(chart, workbook)
    if category_axes and _is_month_axis(chart_spec, labels):
        long_labels = [label for label in labels if len(label) > MAX_MONTH_LABEL]
        if long_labels:
            add("chart_month_labels", f"Month labels longer than {MAX_MONTH_LABEL} characters "
                f"({', '.join(long_labels[:4])}{'...' if len(long_labels) > 4 else ''}) rotate, crowd or clip",
                "Use Jan-Dec as the category labels")
        if any(a.label_rotation != 0 for a in category_axes):
            add("chart_month_labels", "Month labels have no fixed 0 rotation, so Excel may tilt them into the legend",
                "Set the category axis text rotation to 0 (txPr bodyPr rot=\"0\")")

    horizontal = chart.bar_dir == "bar"
    if chart.legend_pos == "b" and category_axes and not horizontal:
        add("chart_legend", "The legend is at the bottom, under the category labels, where they collide",
            "Put the legend at the top (or right) of the chart")

    want = (chart_spec or {}).get("value_axis", {})
    for axis in value_axes:
        if axis.num_fmt is None or axis.num_fmt_linked:
            add("chart_number_format", "The value axis has no number format of its own (it follows the source "
                "cells, so an empty workbook shows 0, 0.2, 0.4 ... or duplicated $ labels)",
                "Set an explicit axis number format, e.g. \"$\"#,##0, with sourceLinked=0")
        elif want.get("num_fmt") and axis.num_fmt != want["num_fmt"]:
            add("chart_number_format", f"The value axis format is {axis.num_fmt}, the spec says {want['num_fmt']}",
                "Use the spec's number format")
        if axis.major_unit is None:
            add("chart_axis_scale", "The value axis has no fixed major unit: on an empty workbook Excel scales it "
                "0-1 and shows duplicate labels ($0 $0 $1 $1)", "Set a fixed major unit (e.g. 1000)")
        elif want.get("major_unit") is not None and axis.major_unit != want["major_unit"]:
            add("chart_axis_scale", f"The value axis major unit is {axis.major_unit:g}, the spec says "
                f"{want['major_unit']:g}", "Use the spec's major unit")
        if axis.min is None and not want.get("allow_negative"):
            add("chart_axis_scale", "The value axis has no fixed minimum", "Set the axis minimum to 0")
        elif want.get("min") is not None and axis.min != want["min"]:
            add("chart_axis_scale", f"The value axis minimum is {axis.min}, the spec says {want['min']}",
                "Use the spec's minimum")
    return problems


def _geometry(charts: list[ChartInfo], sheet, layout: dict, file_name: str, dpi: int) -> list[ChartProblem]:
    problems: list[ChartProblem] = []
    note = f" (at {dpi} DPI, {dpi * 100 // 96}% Windows scaling)"
    placed = [c for c in charts if c.anchor is not None]
    for i, first in enumerate(placed):
        for second in placed[i + 1:]:
            if first.anchor.overlaps(second.anchor):
                problems.append(ChartProblem(
                    "chart_overlap", f"{file_name}:{sheet.name}",
                    f"Chart {first.index} '{first.title}' overlaps chart {second.index} '{second.title}'{note}",
                    "Move or resize the charts so their areas don't overlap"))
    if "print_area" in layout:  # the intended print area, from the spec
        intended = layout["print_area"]
        actual = sheet.print_area_ref.split("!")[-1].replace("$", "") if sheet.print_area_ref else None
        if actual != intended:
            problems.append(ChartProblem(
                "chart_print_area", f"{file_name}:{sheet.name}",
                f"The sheet's print area is {actual or 'not set'}, but the intended print area is {intended}, "
                "so the charts can print across pages", f"Set the print area to {intended}"))
        area = sheet.cell_rect(intended)
        for chart in placed:
            if not chart.anchor.inside(area):
                problems.append(ChartProblem(
                    "chart_print_area", _where(file_name, chart),
                    f"The chart extends past the intended print area {intended}{note}",
                    "Move the chart inside the print area, or widen the area"))
    limit = layout.get("max_right_px")
    for chart in placed:
        if limit is not None and chart.anchor.right > limit:
            problems.append(ChartProblem(
                "chart_width", _where(file_name, chart),
                f"The chart's right edge is at {chart.anchor.right:.0f} px, past the {limit} px that fits the "
                f"screen at 100% zoom{note}", "Make the charts narrower or move them left"))
    return problems


def chart_checks(file_spec: dict, path: Path) -> list[ChartProblem]:
    """Every problem with every chart in one workbook, judged against the file's spec."""
    sheet_specs = {s["name"]: s for s in file_spec.get("sheets", [])}
    workbook = load_workbook(path, read_only=False)
    problems: list[ChartProblem] = []
    seen: set[tuple[str, str]] = set()
    for dpi in CHECK_DPIS:
        for name, sheet in read_charts(path, dpi).items():
            sheet_spec = sheet_specs.get(name, {})
            by_title = {c["title"]: c for c in sheet_spec.get("charts", [])}
            found = []
            if dpi == CHECK_DPIS[0]:
                for chart in sheet.charts:
                    found += _check_chart(chart, by_title.get(chart.title), file_spec["name"], workbook)
            found += _geometry(sheet.charts, sheet, sheet_spec.get("chart_layout", {}), file_spec["name"], dpi)
            for problem in found:
                key = (problem.check, problem.where, re.sub(r"[\d.]+", "#", problem.issue.split(" (at ")[0]))
                if key not in seen:  # one report per problem, not one per DPI (the first DPI's numbers)
                    seen.add(key)
                    problems.append(problem)
    return problems
