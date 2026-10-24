"""Does each boxed cell's text fit inside its box? (QA, deterministic)

A cell drawn as a box (a border or a fill) whose text is wider than the column spills out past the
box in Excel and in the listing pictures, which looks broken: "Jordan & Sam (example)" on the
wedding Dashboard ran past its 16-wide cell. Plain unboxed labels may spill into empty neighbours
on purpose (titles do), so only boxed cells are checked.

Text width is estimated in Excel column-width units (one '0' of the default 11 pt font) from
per-character Calibri widths, the cell's font size and bold.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DEFAULT_WIDTH = 8.43
# Calibri 11 pt widths relative to the digit '0' (7 px), which is what Excel's column width counts.
_NARROW = set("iljI.,:;'!|")
_SEMI = set("frt()[]-/ ")
_WIDE = set("mwMW@%&")


def text_width(text: str, size: float = 11, bold: bool = False) -> float:
    units = 0.0
    for ch in text:
        if ch in _NARROW:
            units += 0.45
        elif ch in _SEMI:
            units += 0.55
        elif ch in _WIDE:
            units += 1.45
        elif ch.isdigit():
            units += 1.0
        elif ch.isupper():
            units += 1.1
        else:
            units += 0.85
    return units * (size / 11) * (1.07 if bold else 1.0)


def _boxed(cell) -> bool:
    border = cell.border
    sides = (border.left, border.right, border.top, border.bottom)
    filled = cell.fill is not None and cell.fill.fill_type not in (None, "none")
    return filled or any(s is not None and s.style for s in sides)


@dataclass
class Overflow:
    sheet: str
    cell: str
    text: str
    needed: float
    width: float


def sheet_overflows(ws, values: dict[tuple[str, str], object] | None = None) -> list[Overflow]:
    """Boxed text cells on one worksheet wider than their column (or merged range)."""
    values = values or {}
    found: list[Overflow] = []
    merged = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = None                    # covered by a merge: not its own box
        merged[(rng.min_row, rng.min_col)] = (rng.min_col, rng.max_col)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or not _boxed(cell):
                continue
            if (cell.row, cell.column) in merged and merged[(cell.row, cell.column)] is None:
                continue
            text = cell.value
            if isinstance(text, str) and text.startswith("="):
                text = values.get((ws.title.upper(), cell.coordinate))
            if not isinstance(text, str) or not text.strip():
                continue  # numbers are checked by the visual render ("####")
            if cell.alignment is not None and (cell.alignment.wrap_text or cell.alignment.shrink_to_fit):
                continue
            first, last = merged.get((cell.row, cell.column)) or (cell.column, cell.column)
            width = sum(ws.column_dimensions[get_column_letter(c)].width or DEFAULT_WIDTH
                        for c in range(first, last + 1))
            font = cell.font
            needed = text_width(text, float(font.sz or 11) if font else 11, bool(font and font.b))
            if needed > width - 0.5:
                found.append(Overflow(ws.title, cell.coordinate, text, round(needed, 1), round(width, 1)))
    return found


def overflows(path: Path, values: dict[tuple[str, str], object] | None = None) -> list[Overflow]:
    """Boxed text cells wider than their box, across a workbook. values: recalculated formula results."""
    wb = load_workbook(path)
    return [o for ws in wb.worksheets for o in sheet_overflows(ws, values)]


def widen_to_fit(ws, cap: float = 60) -> list[Overflow]:
    """Builder helper: widen single columns so typed (not formula) boxed text fits. Returns what it fixed."""
    fixed = []
    for o in sheet_overflows(ws):
        cell = ws[o.cell]
        if any(cell.coordinate in rng for rng in ws.merged_cells.ranges):
            continue  # a merged box: the spec decides how wide it spans
        letter = cell.column_letter
        current = ws.column_dimensions[letter].width or DEFAULT_WIDTH
        ws.column_dimensions[letter].width = max(current, min(cap, round(o.needed + 1.5, 1)))
        fixed.append(o)
    return fixed
