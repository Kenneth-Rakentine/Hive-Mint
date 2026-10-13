"""openpyxl helpers for buyer-facing workbooks: consistent headers, formats, widths,
frozen header rows, input/formula shading and data validation.

Formulas stick to functions that work in both Excel and Google Sheets.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

CURRENCY = '"$"#,##0.00'
RATE = '"$"#,##0.000'
DATE = "mm/dd/yyyy"
NUMBER = "#,##0.0"
INTEGER = "0"
PERCENT = "0.0%"

BRAND = "1F4E5A"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=BRAND)
SUBTITLE_FONT = Font(bold=True, size=12, color=BRAND)
LABEL_FONT = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")  # cells the buyer types into
FORMULA_FILL = PatternFill("solid", fgColor="EEF2F3")  # calculated; leave alone
THIN = Side(style="thin", color="B7C4C8")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
MIN_DATE_SERIAL = "36526"  # 2000-01-01


@dataclass
class Column:
    """One column of a data-entry table.

    formula: template with {r} for the row number and {col[Header]} for another column's letter,
        or a callable (r, letters) -> formula text.
    validation: "date", "money", "number", "percent", "yesno", or "list:<range>".
    """

    header: str
    width: float = 14
    number_format: str | None = None
    formula: str | Callable[[int, dict[str, str]], str] | None = None
    validation: str | None = None


@dataclass
class Table:
    sheet: str
    columns: list[Column]
    rows: int
    letters: dict[str, str] = field(default_factory=dict)

    def col(self, header: str) -> str:
        return self.letters[header]

    def rng(self, header: str) -> str:
        """Absolute data range for a column, sheet-qualified, e.g. Income!$H$2:$H$501."""
        letter = self.letters[header]
        return f"{quote_sheetname(self.sheet)}!${letter}$2:${letter}${self.rows + 1}"


def write_table(ws: Worksheet, columns: list[Column], rows: int) -> Table:
    """Header in row 1, formulas pre-filled for `rows` data rows, frozen header and a filter."""
    table = Table(ws.title, columns, rows)
    for index, column in enumerate(columns, start=1):
        table.letters[column.header] = get_column_letter(index)
    last_row = rows + 1
    for index, column in enumerate(columns, start=1):
        letter = table.letters[column.header]
        header = ws.cell(row=1, column=index, value=column.header)
        header.fill, header.font, header.border = HEADER_FILL, HEADER_FONT, BOX
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[letter].width = column.width
        for r in range(2, last_row + 1):
            cell = ws.cell(row=r, column=index)
            if column.formula:
                text = (column.formula(r, table.letters) if callable(column.formula)
                        else column.formula.format(r=r, col=table.letters))
                cell.value = "=" + text
                cell.fill = FORMULA_FILL
            if column.number_format:
                cell.number_format = column.number_format
        if column.validation:
            add_validation(ws, column.validation, f"{letter}2:{letter}{last_row}")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    return table


def add_validation(ws: Worksheet, kind: str, cells: str) -> DataValidation:
    if kind == "date":
        dv = DataValidation(type="date", operator="greaterThan", formula1=MIN_DATE_SERIAL,
                            error="Enter a date, e.g. 3/15/2026.", errorTitle="Date needed")
    elif kind == "money":
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            error="Enter an amount of 0 or more, without a $ sign.", errorTitle="Amount needed")
    elif kind == "number":
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            error="Enter a number of 0 or more.", errorTitle="Number needed")
    elif kind == "percent":
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                            error="Enter a percentage from 0% to 100%, e.g. 25%.", errorTitle="Percentage needed")
    elif kind == "yesno":
        dv = DataValidation(type="list", formula1='"Yes,No"', error="Choose Yes or No.", errorTitle="Yes or No")
    elif kind.startswith("list:"):
        dv = DataValidation(type="list", formula1="=" + kind.removeprefix("list:"),
                            error="Pick a value from the list (edit the Lists sheet to change it).",
                            errorTitle="Pick from the list")
    else:
        raise ValueError(f"Unknown validation kind: {kind}")
    dv.showErrorMessage = True
    dv.allow_blank = True
    ws.add_data_validation(dv)
    dv.add(cells)
    return dv


def write_how_to(wb: Workbook, title: str, intro: str, steps: list[str], notes: list[str]) -> Worksheet:
    """A short first sheet telling the buyer how to use the workbook."""
    ws = wb.active
    ws.title = "How to Use"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 100
    ws["B1"] = title
    ws["B1"].font = TITLE_FONT
    ws["B2"] = intro
    ws["B2"].alignment = WRAP
    row = 4
    ws.cell(row=row, column=2, value="How to Use").font = SUBTITLE_FONT
    for number, step in enumerate(steps, start=1):
        row += 1
        ws.cell(row=row, column=1, value=number).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=row, column=2, value=step).alignment = WRAP
    row += 2
    ws.cell(row=row, column=2, value="Good to know").font = SUBTITLE_FONT
    for note in notes:
        row += 1
        ws.cell(row=row, column=2, value=note).alignment = WRAP
    row += 2
    key = ws.cell(row=row, column=2, value="White table cells and yellow cells are for your entries. Grey cells "
                                           "calculate automatically; don't type over them.")
    key.alignment = WRAP
    key.fill = INPUT_FILL
    ws.sheet_view.showGridLines = False
    return ws


def write_list_sheet(wb: Workbook, lists: dict[str, list[str]], length: int = 30) -> dict[str, str]:
    """A 'Lists' sheet of editable dropdown values. Returns header -> absolute range."""
    ws = wb.create_sheet("Lists")
    ranges = {}
    for index, (header, values) in enumerate(lists.items(), start=1):
        letter = get_column_letter(index)
        cell = ws.cell(row=1, column=index, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[letter].width = max(24, len(header) + 4)
        for offset in range(length):
            value_cell = ws.cell(row=2 + offset, column=index, value=values[offset] if offset < len(values) else None)
            value_cell.fill = INPUT_FILL
        ranges[header] = f"Lists!${letter}$2:${letter}${length + 1}"
    ws.freeze_panes = "A2"
    return ranges


def label(ws: Worksheet, cell: str, text: str) -> None:
    ws[cell] = text
    ws[cell].font = LABEL_FONT


def value_cell(ws: Worksheet, cell: str, value, number_format: str | None = None, *, entry: bool = False) -> None:
    ws[cell] = value
    ws[cell].fill = INPUT_FILL if entry else FORMULA_FILL
    ws[cell].border = BOX
    if number_format:
        ws[cell].number_format = number_format


def header_row(ws: Worksheet, row: int, first_column: int, headers: list[str]) -> None:
    for offset, text in enumerate(headers):
        cell = ws.cell(row=row, column=first_column + offset, value=text)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BOX


def year_criteria(rng: str, year_cell: str) -> str:
    """SUMIFS/COUNTIFS criteria pair limiting a date range to the tax year in year_cell."""
    return f'{rng},">="&DATE({year_cell},1,1),{rng},"<"&DATE({year_cell}+1,1,1)'
