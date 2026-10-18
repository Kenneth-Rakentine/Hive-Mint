"""Compliance department: deterministic checks on what the product says. No AI.

- Claims: none of the spec's claims_not_to_make phrases may appear in any file (case-insensitive).
- Disclaimer: a spec with a regulated_topic (tax, legal, financial) needs "not <topic> advice" in the
  README and the PDF.
- Brands: brand or trademark names in the files (other than the software a file opens in).

A claim found is a clear failure: back to REWORK with a brief. A missing disclaimer or a brand
name is ambiguous (context matters), so it goes to HUMAN_REVIEW with the reason; CLAUDE.md:
compliance ambiguity goes to a person, never a guess. A clean product moves on to SEO.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import object_session

from agents.departments import StageResult
from agents.rework import create_brief
from app.lifecycle import log_event
from app.models import Build, Job, Product
from app.states import Outcome, ProductStatus
from config.settings import Settings
from tools.brands import brand_hits
from tools.digest import pdf_text
from tools.spec_store import product_spec

S = ProductStatus
ACTOR = "compliance"


def product_texts(spec: dict, build_dir: Path) -> dict[str, str]:
    """Every buyer-visible string, by file name."""
    texts = {}
    for file_spec in spec["files"]:
        path = build_dir / file_spec["name"]
        if not path.is_file():
            continue
        if file_spec["kind"] == "xlsx":
            wb = load_workbook(path)
            texts[file_spec["name"]] = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row
                                                if isinstance(c.value, str) and not c.value.startswith("="))
        elif file_spec["kind"] == "pdf":
            texts[file_spec["name"]] = pdf_text(path)
        else:
            texts[file_spec["name"]] = path.read_text(encoding="utf-8", errors="replace")
    return texts


def compliance_findings(spec: dict, texts: dict[str, str]) -> tuple[list[str], list[str]]:
    """(claim violations, ambiguities for a person)."""
    claims, ambiguous = [], []
    for name, text in texts.items():
        low = text.lower()
        claims += [f"{name} says \"{c}\"" for c in spec.get("claims_not_to_make", []) if c.lower() in low]
        if hits := brand_hits(text, body_text=True):
            ambiguous.append(f"{name} mentions brand names: {', '.join(hits)}")
    if topic := spec.get("regulated_topic"):
        phrase = f"not {topic} advice"
        for file_spec in spec["files"]:
            if file_spec["kind"] in ("pdf", "text") and phrase not in " ".join(
                    texts.get(file_spec["name"], "").lower().split()):
                ambiguous.append(f"{file_spec['name']} lacks the disclaimer \"{phrase}\"")
    return claims, ambiguous


def compliance_gate(settings: Settings):
    def handler(product: Product, job: Job) -> StageResult:
        session = object_session(product)
        spec = product_spec(session, product, settings.products_dir)
        build = session.scalar(select(Build).where(Build.product_id == product.id)
                               .order_by(Build.version.desc()).limit(1))
        if spec is None or build is None:
            return StageResult(Outcome.AMBIGUOUS, S.HUMAN_REVIEW, "Compliance: no spec or build to check")
        claims, ambiguous = compliance_findings(spec, product_texts(spec, Path(build.path)))
        data = {"build_id": build.id, "claims": claims, "ambiguous": ambiguous}
        log_event(session, "compliance.checked", actor=ACTOR, product=product, job=job,
                  message=f"{len(claims)} claim(s), {len(ambiguous)} ambiguity(ies)", data=data)
        if claims:
            brief = create_brief(session, product, source="COMPLIANCE", actor=ACTOR,
                                 summary="Product text makes claims the spec forbids",
                                 required_changes=[{"id": f"remove_forbidden_claims_v{build.version}",
                                                    "change": "Remove or reword: " + "; ".join(claims),
                                                    "reason": "Forbidden claims mislead buyers and create risk",
                                                    "acceptance": "No claims_not_to_make phrase appears in any file"}],
                                 findings=[{"claim": c} for c in claims], job=job, build=build)
            return StageResult(Outcome.FAIL, S.REWORK, "Compliance: " + "; ".join(claims),
                               data | {"rework_brief_id": brief.id})
        if ambiguous:
            return StageResult(Outcome.AMBIGUOUS, S.HUMAN_REVIEW,
                               "Compliance needs a human decision: " + "; ".join(ambiguous), data)
        return StageResult(Outcome.PASS, S.SEO, f"Compliance passed v{build.version}: no forbidden claims, "
                                                "disclaimers present, no brand names", data)

    handler.provider = "deterministic"
    return handler
