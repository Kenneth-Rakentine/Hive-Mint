"""Testing department: "does it function?" Deterministic checks only, no AI.

Every check is driven by the product spec (config/specs/<slug>.json): which files must
exist, which sheets and columns each workbook needs, which scenarios to recalculate and
what the answers must be. Results are stored as TestResult rows linked to the product,
job, agent run and build. Any failure sends the product to REWORK with the reasons; a
clean run moves it to QA.
"""

from __future__ import annotations

import re
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import object_session

from agents.departments import StageResult
from app.models import AgentRun, Build, Job, Product, TestResult
from app.states import Outcome, ProductStatus
from config.settings import Settings
from tools.packaging import sha256
from tools.recalc import RecalcError, Values, choose_engine, recalculate
from tools.spec import excel_serial, excel_value, file_names, load_spec, sheet_spec

S = ProductStatus
TOLERANCE = 0.005  # half a cent
MAX_REASONS_IN_SUMMARY = 8
_STRING_LITERAL = re.compile(r'"[^"]*"')
_CHART_REF = re.compile(r"^(?:'(?P<quoted>(?:[^']|'')+)'|(?P<plain>[^!']+))!"
                        r"(?P<range>\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)$")
_FUNCTION = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-z][A-Za-z0-9_.]*)\s*\(")


@dataclass(frozen=True)
class Check:
    check: str
    target: str
    passed: bool
    detail: str = ""

    def reason(self) -> str:
        return f"{self.check} [{self.target}]: {self.detail}"


class Testing:
    __test__ = False  # not a pytest test class
    provider = "deterministic"

    def __init__(self, settings: Settings):
        self.settings = settings

    def __call__(self, product: Product, job: Job) -> StageResult:
        session = object_session(product)
        spec = load_spec(product.slug)
        if spec is None:
            return StageResult(Outcome.FAIL, S.REWORK, f"Testing: no specification for '{product.slug}'")
        build = session.scalar(select(Build).where(Build.product_id == product.id)
                               .order_by(Build.version.desc()).limit(1))
        if build is None:
            return StageResult(Outcome.FAIL, S.REWORK, "Testing: no build recorded for this product")

        engine, soffice = choose_engine(self.settings.recalc_engine, self.settings.soffice_path)
        checks = run_checks(spec, Path(build.path), engine=engine, soffice=soffice,
                            expected_zip_sha256=build.zip_sha256)
        run = session.scalar(select(AgentRun).where(AgentRun.job_id == job.id).order_by(AgentRun.id.desc()).limit(1))
        session.add_all(TestResult(product_id=product.id, job_id=job.id, agent_run_id=run.id if run else None,
                                   build_id=build.id, check=c.check, target=c.target, passed=c.passed,
                                   detail=c.detail) for c in checks)

        failed = [c for c in checks if not c.passed]
        data = {"build_id": build.id, "version": build.version, "engine": engine, "checks": len(checks),
                "failed": len(failed)}
        if failed:
            reasons = "; ".join(c.reason() for c in failed[:MAX_REASONS_IN_SUMMARY])
            more = f"; and {len(failed) - MAX_REASONS_IN_SUMMARY} more" if len(failed) > MAX_REASONS_IN_SUMMARY else ""
            return StageResult(Outcome.FAIL, S.REWORK,
                               f"Testing failed {len(failed)} of {len(checks)} checks on v{build.version}: "
                               f"{reasons}{more}", data | {"failures": [c.reason() for c in failed]})
        return StageResult(Outcome.PASS, S.QA,
                           f"Testing passed all {len(checks)} checks on v{build.version} (recalc: {engine})", data)


# --------------------------------------------------------------------------- checks

def run_checks(spec: dict, build_dir: Path, *, engine: str = "formulas", soffice: str | None = None,
               expected_zip_sha256: str | None = None) -> list[Check]:
    checks = check_files(spec, build_dir)
    checks += check_zip(spec, build_dir, expected_zip_sha256)
    rules = spec.get("spreadsheet_rules", {})
    for file_spec in spec["files"]:
        path = build_dir / file_spec["name"]
        if not path.is_file():
            continue  # already reported by check_files
        if file_spec["kind"] == "xlsx":
            checks += check_workbook(spec, file_spec, rules, path, engine=engine, soffice=soffice)
        elif file_spec["kind"] == "pdf":
            checks += check_pdf(file_spec, path)
        elif file_spec["kind"] == "text":
            checks += check_text(file_spec, path)
    return checks


def check_files(spec: dict, build_dir: Path) -> list[Check]:
    checks = []
    for name in file_names(spec):
        path = build_dir / name
        ok = path.is_file() and path.stat().st_size > 0
        checks.append(Check("file_exists", name, ok, "" if ok else "missing or empty"))
    expected = set(file_names(spec)) | {spec["zip_name"]}
    present = {p.name for p in build_dir.iterdir()} if build_dir.is_dir() else set()
    extra = sorted(present - expected)
    checks.append(Check("build_folder_clean", build_dir.name, not extra,
                        f"unexpected files: {', '.join(extra)}" if extra else ""))
    return checks


def check_zip(spec: dict, build_dir: Path, expected_sha256: str | None) -> list[Check]:
    name = spec["zip_name"]
    path = build_dir / name
    if not path.is_file():
        return [Check("zip_opens", name, False, "ZIP is missing")]
    checks = []
    if expected_sha256:
        actual = sha256(path)
        checks.append(Check("zip_integrity", name, actual == expected_sha256,
                            "" if actual == expected_sha256 else "ZIP changed since Production recorded it"))
    try:
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            members = archive.namelist()
            contents = {m: archive.read(m) for m in members if not m.endswith("/")}
    except (zipfile.BadZipFile, OSError) as error:
        return checks + [Check("zip_opens", name, False, f"{type(error).__name__}: {error}")]
    checks.append(Check("zip_opens", name, bad is None, f"corrupt member: {bad}" if bad else ""))

    expected = file_names(spec)
    missing = sorted(set(expected) - set(members))
    extra = sorted(set(members) - set(expected))
    dupes = sorted({m for m in members if members.count(m) > 1})
    problems = ([f"missing {', '.join(missing)}"] if missing else []) + \
               ([f"unexpected {', '.join(extra)}"] if extra else []) + \
               ([f"duplicated {', '.join(dupes)}"] if dupes else [])
    checks.append(Check("zip_contents", name, not problems, "; ".join(problems)))

    stale = [m for m in expected if m in contents and (build_dir / m).is_file()
             and contents[m] != (build_dir / m).read_bytes()]
    checks.append(Check("zip_matches_build", name, not stale,
                        f"differs from build folder: {', '.join(stale)}" if stale else ""))
    return checks


# --------------------------------------------------------------------------- workbooks

def check_workbook(spec: dict, file_spec: dict, rules: dict, path: Path, *, engine: str = "formulas",
                   soffice: str | None = None) -> list[Check]:
    name = file_spec["name"]
    try:
        wb = load_workbook(path)
    except Exception as error:  # any failure to open is the finding
        return [Check("workbook_opens", name, False, f"{type(error).__name__}: {error}")]
    checks = [Check("workbook_opens", name, True)]

    missing_sheets = [s["name"] for s in file_spec["sheets"] if s["name"] not in wb.sheetnames]
    checks.append(Check("sheets_present", name, not missing_sheets,
                        f"missing sheets: {', '.join(missing_sheets)}" if missing_sheets else ""))
    for sheet in file_spec["sheets"]:
        if sheet["name"] in wb.sheetnames and "columns" in sheet:
            checks += _check_table(wb[sheet["name"]], sheet, f"{name}:{sheet['name']}")
        if sheet["name"] in wb.sheetnames and "charts" in sheet:
            checks.append(_check_charts(wb, sheet, f"{name}:{sheet['name']}"))

    checks += _check_formulas(wb, rules, name)
    if rate_table := file_spec.get("rate_table"):
        checks.append(_check_rate_table(wb, rate_table, spec["parameters"][rate_table["parameter"]], name))
    if missing_sheets:
        return checks  # scenarios address those sheets; their failures would only repeat this one

    forbidden = set(rules.get("forbidden_values", []))
    with tempfile.TemporaryDirectory(prefix="hivemint-test-") as tmp:
        workdir = Path(tmp)
        template = _recalc(path, engine, workdir / "template", soffice)
        if isinstance(template, Check):
            return checks + [template]
        checks.append(_error_scan(template, forbidden, f"{name}: empty template"))
        for index, scenario in enumerate(file_spec.get("scenarios", [])):
            target = f"{name}: {scenario['name']}"
            if not scenario.get("inputs"):
                values = template
            else:
                filled = workdir / f"scenario{index}" / path.name
                problem = _write_inputs(path, file_spec, scenario["inputs"], filled)
                if problem:
                    checks.append(Check("scenario", target, False, problem))
                    continue
                values = _recalc(filled, engine, workdir / f"scenario{index}", soffice)
                if isinstance(values, Check):
                    checks.append(values)
                    continue
                checks.append(_error_scan(values, forbidden, target))
            checks.append(_check_expectations(wb, file_spec, scenario, values, target))
    return checks


def _check_table(ws, sheet: dict, target: str) -> list[Check]:
    header_row = sheet.get("header_row", 1)
    headers = _headers(ws, header_row)
    missing = [c for c in sheet["columns"] if c not in headers]
    checks = [Check("columns_present", target, not missing, f"missing columns: {', '.join(missing)}" if missing else "")]
    if sheet.get("freeze_header"):
        frozen_row = _frozen_rows(ws)
        ok = frozen_row >= header_row
        checks.append(Check("header_frozen", target, ok, "" if ok else f"rows frozen: {frozen_row}"))
    unvalidated = []
    for column in sheet.get("validated_columns", []):
        if column not in headers:
            continue  # reported as missing above
        cell = f"{headers[column]}{header_row + 1}"
        if not any(_covers(dv.sqref, cell) for dv in ws.data_validations.dataValidation):
            unvalidated.append(column)
    if sheet.get("validated_columns"):
        checks.append(Check("data_validation", target, not unvalidated,
                            f"no validation on: {', '.join(unvalidated)}" if unvalidated else ""))
    return checks


def _headers(ws, header_row: int) -> dict[str, str]:
    """Header text -> column letter."""
    return {str(c.value).strip(): c.column_letter for c in ws[header_row] if c.value is not None}


def _frozen_rows(ws) -> int:
    if not ws.freeze_panes:
        return 0
    return int(re.sub(r"[A-Z]+", "", ws.freeze_panes)) - 1


def _covers(sqref, cell: str) -> bool:
    col, row = column_index_from_string(re.sub(r"\d+", "", cell)), int(re.sub(r"[A-Z]+", "", cell))
    for rng in str(sqref).split():
        min_col, min_row, max_col, max_row = range_boundaries(rng)
        if min_col <= col <= max_col and min_row <= row <= max_row:
            return True
    return False


def _check_formulas(wb, rules: dict, name: str) -> list[Check]:
    allowed = set(rules.get("allowed_functions", []))
    forbidden = rules.get("forbidden_values", [])
    unknown: dict[str, str] = {}
    broken: list[str] = []
    external: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                where = f"{ws.title}!{cell.coordinate}"
                code = _STRING_LITERAL.sub('""', cell.value)
                for fn in _FUNCTION.findall(code):
                    if allowed and fn.upper() not in allowed:
                        unknown.setdefault(fn.upper(), where)
                if any(err in code for err in forbidden):
                    broken.append(where)
                if "[" in code:
                    external.append(where)
    checks = []
    if allowed:
        checks.append(Check("functions_compatible", name, not unknown,
                            "not allowed in Excel + Google Sheets: " +
                            ", ".join(f"{fn} (first at {at})" for fn, at in sorted(unknown.items()))
                            if unknown else ""))
    checks.append(Check("formula_text_clean", name, not (broken or external),
                        "; ".join(filter(None, [
                            f"error literal in {', '.join(broken[:5])}" if broken else "",
                            f"external link in {', '.join(external[:5])}" if external else "",
                        ]))))
    return checks


def _chart_title(chart) -> str:
    try:
        return "".join(run.t for p in chart.title.tx.rich.p for run in (p.r or []))
    except AttributeError:
        return ""


def _series_refs(series) -> list[str]:
    refs = []
    for source in (series.val, series.cat):
        if source is None:
            continue
        for ref in (getattr(source, "numRef", None), getattr(source, "strRef", None)):
            if ref is not None and ref.f:
                refs.append(ref.f)
    if series.tx is not None and series.tx.strRef is not None:
        refs.append(series.tx.strRef.f)
    return refs


def _ref_problem(wb, ref: str) -> str | None:
    """A chart data reference must name an existing sheet and a range with formulas or values in it."""
    match = _CHART_REF.match(ref)
    if not match:
        return f"unreadable reference {ref}"
    sheet = match["quoted"].replace("''", "'") if match["quoted"] else match["plain"]
    if sheet not in wb.sheetnames:
        return f"{ref} points at missing sheet '{sheet}'"
    ws = wb[sheet]
    try:
        min_col, min_row, max_col, max_row = range_boundaries(match["range"].replace("$", ""))
    except ValueError:
        return f"unreadable range in {ref}"
    cells = [ws.cell(r, c).value for r in range(min_row, (max_row or min_row) + 1)
             for c in range(min_col, (max_col or min_col) + 1)]
    if all(v is None for v in cells):
        return f"{ref} is empty"
    return None


def _check_charts(wb, sheet: dict, target: str) -> Check:
    """Charts exist with the spec'd titles and series, and every series reads real cells."""
    charts = wb[sheet["name"]]._charts
    expected = sheet["charts"]
    problems = []
    if len(charts) != len(expected):
        problems.append(f"{len(charts)} chart(s), expected {len(expected)}")
    by_title = {_chart_title(c): c for c in charts}
    for want in expected:
        chart = by_title.get(want["title"])
        if chart is None:
            problems.append(f"no chart titled '{want['title']}'")
            continue
        if len(chart.series) != want.get("series", len(chart.series)):
            problems.append(f"'{want['title']}' has {len(chart.series)} series, expected {want['series']}")
        for series in chart.series:
            refs = _series_refs(series)
            if not refs:
                problems.append(f"'{want['title']}' has a series with no data reference")
            problems += [f"'{want['title']}': {p}" for ref in refs if (p := _ref_problem(wb, ref))]
    return Check("charts", target, not problems, "; ".join(problems))


def _check_rate_table(wb, rate_table: dict, expected: list[dict], name: str) -> Check:
    target = f"{name}:{rate_table['sheet']}"
    if rate_table["sheet"] not in wb.sheetnames:
        return Check("rate_table", target, False, "sheet missing")
    ws = wb[rate_table["sheet"]]
    headers = _headers(ws, 1)
    if rate_table["date_column"] not in headers or rate_table["rate_column"] not in headers:
        return Check("rate_table", target, False, "rate columns missing")
    actual = []
    for row in range(2, ws.max_row + 1):
        when = ws[f"{headers[rate_table['date_column']]}{row}"].value
        rate = ws[f"{headers[rate_table['rate_column']]}{row}"].value
        if when is None and rate is None:
            continue
        actual.append((when.date().isoformat() if hasattr(when, "date") else str(when), rate))
    wanted = [(r["effective_from"], r["rate"]) for r in expected]
    ok = len(actual) == len(wanted) and all(a[0] == w[0] and _close(a[1], w[1]) for a, w in zip(actual, wanted))
    return Check("rate_table", target, ok, "" if ok else f"expected {wanted}, found {actual}")


def _recalc(path: Path, engine: str, workdir: Path, soffice: str | None) -> Values | Check:
    workdir.mkdir(parents=True, exist_ok=True)
    try:
        return recalculate(path, engine, workdir, soffice)
    except RecalcError as error:
        return Check("recalculates", path.name, False, str(error))


def _error_scan(values: Values, forbidden: set[str], target: str) -> Check:
    errors = [f"{sheet.title()}!{cell}={value}" for (sheet, cell), value in sorted(values.items())
              if isinstance(value, str) and value in forbidden]
    detail = ""
    if errors:
        detail = f"{len(errors)} error cell(s): {', '.join(errors[:5])}" + (" ..." if len(errors) > 5 else "")
    return Check("no_error_values", target, not errors, detail)


def _resolve(wb, file_spec: dict, ref: dict) -> str:
    """Spec reference -> cell coordinate. Raises LookupError with a readable reason."""
    ws = wb[ref["sheet"]]
    if "label" in ref:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() == ref["label"]:
                    col = cell.column + ref.get("offset", 1)
                    return f"{get_column_letter(col)}{cell.row + ref.get('row_offset', 0)}"
        raise LookupError(f"label '{ref['label']}' not found on {ref['sheet']}")
    sheet = sheet_spec(file_spec, ref["sheet"])
    header_row = sheet.get("header_row", 1)
    headers = _headers(ws, header_row)
    if ref["column"] not in headers:
        raise LookupError(f"column '{ref['column']}' not found on {ref['sheet']}")
    return f"{headers[ref['column']]}{header_row + ref['row']}"


def _describe(ref: dict) -> str:
    if "label" in ref:
        offsets = "".join(f" {k}={ref[k]}" for k in ("row_offset", "offset") if k in ref)
        return f"{ref['sheet']} '{ref['label']}'{offsets}"
    return f"{ref['sheet']} row {ref['row']} '{ref['column']}'"


def _write_inputs(path: Path, file_spec: dict, inputs: list[dict], out: Path) -> str | None:
    """Copy the workbook with scenario inputs typed in. Returns a problem description or None."""
    wb = load_workbook(path)
    try:
        for ref in inputs:
            wb[ref["sheet"]][_resolve(wb, file_spec, ref)] = excel_value(ref["value"])
    except LookupError as error:
        return f"cannot place input: {error}"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return None


def _check_expectations(wb, file_spec: dict, scenario: dict, values: Values, target: str) -> Check:
    misses = []
    for ref in scenario.get("expect", []):
        try:
            cell = _resolve(wb, file_spec, ref)
        except LookupError as error:
            misses.append(str(error))
            continue
        actual = values.get((ref["sheet"].upper(), cell))
        expected = excel_value(ref["value"])
        if hasattr(expected, "isoformat"):
            expected = excel_serial(expected)
        if not _matches(actual, expected):
            misses.append(f"{_describe(ref)} ({ref['sheet']}!{cell}) expected {expected!r}, got {actual!r}")
    return Check("scenario", target, not misses, "; ".join(misses))


def _matches(actual, expected) -> bool:
    if expected == "":
        return actual in ("", None)
    if isinstance(expected, int | float) and not isinstance(expected, bool):
        return isinstance(actual, int | float) and not isinstance(actual, bool) and _close(actual, expected)
    return actual == expected


def _close(a, b) -> bool:
    try:
        return abs(float(a) - float(b)) <= TOLERANCE
    except (TypeError, ValueError):
        return False


# --------------------------------------------------------------------------- PDF and text

def _normalize(text: str) -> str:
    text = text.replace("’", "'").replace("‘", "'").replace("&amp;", "&")
    return re.sub(r"\s+", " ", text).strip()


def check_pdf(file_spec: dict, path: Path) -> list[Check]:
    name = file_spec["name"]
    try:
        reader = PdfReader(path)
        pages = len(reader.pages)
        text = _normalize(" ".join(page.extract_text() or "" for page in reader.pages))
    except Exception as error:  # any failure to open or read is the finding
        return [Check("pdf_opens", name, False, f"{type(error).__name__}: {error}")]
    low, high = file_spec.get("min_pages", 1), file_spec.get("max_pages", 1000)
    missing = [h for h in file_spec.get("required_headings", []) if _normalize(h) not in text]
    return [
        Check("pdf_opens", name, True),
        Check("pdf_page_count", name, low <= pages <= high, f"{pages} pages (expected {low}-{high})"),
        Check("pdf_headings", name, not missing, f"missing headings: {', '.join(missing)}" if missing else ""),
    ]


def check_text(file_spec: dict, path: Path) -> list[Check]:
    name = file_spec["name"]
    try:
        text = path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as error:
        return [Check("text_readable", name, False, f"{type(error).__name__}: {error}")]
    missing = [p for p in file_spec.get("required_phrases", []) if p not in text]
    return [
        Check("text_readable", name, True),
        Check("text_phrases", name, not missing, f"missing: {', '.join(missing)}" if missing else ""),
    ]
