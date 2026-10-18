"""QA department: "is it good?" Deterministic checks first, then local-model passes.

Deterministic (no AI): spelling against a dictionary plus a domain allowlist, consistent
sheet and column naming, every spec feature present in the build, the README and PDF
naming every file in the download, and chart checks read from the workbook XML
(tools/chart_checks.py: title overlay, month labels, legend placement, axis format and
scale, chart overlap, print area and screen width).

Then the Tier 1 local model (Ollama) reads a compact text digest of the product, never raw
files, and reports clarity problems. Visual QA renders each workbook the way a buyer sees it
(tools/render.py: Excel, else LibreOffice, else a clear skip) and the same local model looks
at every chart PNG. A 4B model judges poorly but transcribes well, so it reports what it
sees (where the title sits, legend overlap, every axis label as printed, rotation, cut-off
text, anything else broken) and code turns that into findings: e.g. a value repeated on the
value axis is a duplicate-label defect, whatever the model thought of it.

The model cannot fail a product on its own: its findings are capped at "major", and their
deductions at QA_MODEL_MAX_DEDUCTION (text) and QA_VISION_MAX_DEDUCTION (charts), which
together may not exceed 100 - QA_PASS_SCORE. A critical or major deterministic finding, or
a score under QA_PASS_SCORE, sends the product to REWORK with a rework brief.
"""

from __future__ import annotations

import json
import re
from collections.abc import Callable
from dataclasses import asdict, dataclass
from functools import lru_cache
from pathlib import Path
from typing import Literal

import httpx
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import object_session

from agents.departments import StageResult
from agents.rework import create_brief
from app.lifecycle import log_event
from app.models import Build, Job, Product, Review
from app.states import Department, Outcome, ProductStatus
from config.settings import ROOT, Settings
from providers.ollama import OllamaProvider
from tools.chart_checks import chart_checks
from tools.digest import pdf_text, product_digest, workbook_text
from tools.render import RENDERS_DIR, RenderResult, render_workbooks
from tools.spec_store import product_spec

S = ProductStatus
SEVERITY_POINTS = {"critical": 40, "major": 15, "minor": 5, "info": 0}
BLOCKING = frozenset({"critical", "major"})
AI_SOURCES = ("model", "vision")
ALLOWLIST = ROOT / "config" / "qa_words.txt"
SMALL_WORDS = frozenset({"a", "an", "and", "as", "at", "by", "for", "in", "of", "on", "or", "per", "the", "to",
                         "vs", "with"})
_WORD = re.compile(r"[A-Za-z][A-Za-z']*")
_SKIP_TOKEN = re.compile(r"\S*[_./@\\]\S*|\S*\d\S*")  # file names, URLs, emails, anything with digits

SYSTEM = ("You are the QA department of a company that sells spreadsheet templates to small-business owners. "
          "You check whether a product is clear and easy for a non-technical buyer to use. You only judge clarity "
          "and instructions: formulas, totals and file integrity were already tested by code. Be specific and "
          "fair; do not invent problems. If the instructions are clear, return few or no findings.")
PROMPT = """Review this product digest (sheet names, column headers, instruction text, chart titles, the PDF
guide and the README). Look for: confusing or contradictory instructions, steps a beginner could not follow,
undefined jargon, places where the buyer would not know where to type, and mismatches between the guide and
the workbook.

Text marked "[... trimmed]" was shortened to fit; do not report the cut itself as a problem.

Return JSON: clarity_score from 0 (unusable) to 100 (perfectly clear), and up to 8 findings, each with
severity "minor" or "major", where (sheet or document), issue, and suggestion.

PRODUCT DIGEST
{digest}
"""


VISION_SYSTEM = ("You inspect one chart image exported from Excel, for a company that sells spreadsheet templates. "
                 "Report exactly what you see. Do not guess and do not invent problems.")
VISION_PROMPT = """This is the chart "{title}" on the {sheet} tab of {state}

Fill in what you observe:
- title_position: "above_plot" if the title sits in its own space above the plot area, "on_plot" if it is
  drawn on top of the plot area, gridlines or axis lines, or "none" if there is no title
- legend_overlaps: true if the legend touches or covers the plot, the gridlines or any axis label
- value_axis_labels: every label on the value (numbers/$) axis, in order, exactly as printed, repeats included
- category_axis_labels: every label on the category axis, in order, exactly as printed
- labels_rotated: true if any axis label is tilted or vertical instead of horizontal
- text_cut_off: true if any text is cut off at an edge or only partly visible
- unreadable_colors: true if any text or bar color has too little contrast to read
- other_problems: anything else that looks broken or unprofessional (up to 3), or an empty list
"""
# What the chart should look like, by the workbook's role in the spec ("template" unless it is an example).
VISION_STATES = {
    "template": "a blank template: no data has been entered yet,\nso the bars are empty. Empty bars are expected. "
                "Everything else (title, legend, axis labels) must still look right.",
    "example": "a filled-in example workbook: sample data has been\nentered, so the bars should show values. "
               "Report empty bars under other_problems. Title, legend and axis labels must look right.",
}
# The model often "reports" the empty template itself; that is expected, not a defect.
_EMPTY_TEMPLATE_NOTE = re.compile(r"\b(no data|blank template|empty (bars|chart|template)|expected)\b", re.I)


@dataclass
class Finding:
    source: str  # "deterministic", "model" (text clarity), "vision" (chart images) or "render"
    check: str
    severity: str
    where: str
    issue: str
    suggestion: str = ""


class ModelFinding(BaseModel):
    severity: Literal["minor", "major"]
    where: str
    issue: str
    suggestion: str = ""


class ModelReview(BaseModel):
    clarity_score: int = Field(ge=0, le=100)
    findings: list[ModelFinding] = Field(default_factory=list, max_length=8)


class OtherProblem(BaseModel):
    severity: Literal["minor", "major"]
    issue: str


class ChartObservation(BaseModel):
    """What the model sees in one chart image. Code, not the model, decides what is a defect."""

    title_position: Literal["above_plot", "on_plot", "none"]
    legend_overlaps: bool
    value_axis_labels: list[str] = Field(default_factory=list, max_length=40)
    category_axis_labels: list[str] = Field(default_factory=list, max_length=40)
    labels_rotated: bool
    text_cut_off: bool
    unreadable_colors: bool = False
    other_problems: list[OtherProblem] = Field(default_factory=list, max_length=3)


def judge_observation(seen: ChartObservation, *, template: bool = True) -> list[tuple[str, str, str, str]]:
    """(severity, problem, issue, suggestion) for each defect the observation shows. On a blank template,
    remarks about missing data are expected and dropped; on an example workbook they are defects."""
    out = []
    if seen.title_position == "on_plot":
        out.append(("major", "overlapping_text", "The title is drawn on top of the plot area",
                    "Turn the title overlay off so the plot starts below it"))
    if seen.legend_overlaps:
        out.append(("major", "overlapping_text", "The legend overlaps the plot or the axis labels",
                    "Move the legend to the top and turn its overlay off"))
    for axis, labels in (("value", seen.value_axis_labels), ("category", seen.category_axis_labels)):
        shown = [label.strip() for label in labels if label.strip()]
        repeated = sorted({label for label in shown if shown.count(label) > 1})
        if repeated:
            out.append(("major", "duplicate_axis_labels",
                        f"The {axis} axis repeats labels ({', '.join(repeated[:4])}): read as {' '.join(shown[:12])}",
                        "Fix the axis scale and number format so every label is different"))
    if seen.labels_rotated:
        out.append(("minor", "unreadable", "Axis labels are rotated", "Use short labels and a fixed 0 rotation"))
    if seen.text_cut_off:
        out.append(("major", "clipped_label", "Text is cut off at an edge", "Give the labels room or shorten them"))
    if seen.unreadable_colors:
        out.append(("minor", "unreadable", "Some text or bar colors have too little contrast",
                    "Use darker text and distinct series colors"))
    out += [(p.severity, "other", p.issue, "") for p in seen.other_problems
            if not (template and _EMPTY_TEMPLATE_NOTE.search(p.issue))]
    return out


# --------------------------------------------------------------------------- deterministic checks

@lru_cache(maxsize=1)
def _spellchecker():
    from spellchecker import SpellChecker

    return SpellChecker()


def allowlist(path: Path = ALLOWLIST) -> set[str]:
    if not path.is_file():
        return set()
    lines = path.read_text(encoding="utf-8").splitlines()
    return {line.strip().lower() for line in lines if line.strip() and not line.startswith("#")}


def words(text: str) -> list[str]:
    text = _SKIP_TOKEN.sub(" ", text)
    out = []
    for word in _WORD.findall(text.replace("-", " ")):
        word = word.strip("'")
        if word.endswith("'s"):
            word = word[:-2]
        if len(word) > 1 and not word.isupper():  # skip acronyms (IRS, PDF, PPE)
            out.append(word.lower())
    return out


def check_spelling(texts: dict[str, str], extra_words: set[str]) -> list[Finding]:
    """texts: where -> text. One finding listing every unknown word, with where it first appears."""
    first_seen: dict[str, str] = {}
    for where, text in texts.items():
        for word in words(text):
            first_seen.setdefault(word, where)
    unknown = sorted(w for w in _spellchecker().unknown(first_seen) if w not in extra_words)
    if not unknown:
        return []
    severity = "major" if len(unknown) >= 5 else "minor"
    listed = ", ".join(f"'{w}' ({first_seen[w]})" for w in unknown[:20])
    return [Finding("deterministic", "spelling", severity, "all text", f"{len(unknown)} unrecognised word(s): {listed}",
                    "Fix the spelling, or add genuine domain words to config/qa_words.txt")]


def _title_case_problem(name: str) -> bool:
    tokens = [t for t in re.split(r"[\s/]+", name) if t and t[0].isalpha()]
    for index, token in enumerate(tokens):
        if index > 0 and token.lower() in SMALL_WORDS:
            if token != token.lower():
                return True
        elif not token[0].isupper():
            return True
    return False


def _naming_key(name: str) -> str:
    """Names differing only in case, spacing or punctuation ("Paid?" vs "Paid", "Set Aside" vs "Set-Aside")
    are the same concept spelled two ways. Singular vs plural is allowed: "IRS Rates" (a tab) holds rates and
    "IRS Rate" (a column) is one rate."""
    return re.sub(r"[^a-z0-9]", "", name.lower())


def check_naming(sheet_names: list[str], headers: dict[str, list[str]]) -> list[Finding]:
    """Title Case for sheet names and headers, and one spelling per concept across sheets."""
    findings = []
    bad_case = [n for n in sheet_names if _title_case_problem(n)]
    bad_case += [f"{sheet}: {h}" for sheet, hs in headers.items() for h in hs if _title_case_problem(h)]
    if bad_case:
        findings.append(Finding("deterministic", "naming", "minor", "sheet and column names",
                                f"Not in Title Case: {', '.join(bad_case[:12])}",
                                "Use Title Case consistently (small words like 'per' and 'to' stay lowercase)"))
    forms: dict[str, set[str]] = {}
    for name in sheet_names + [h for hs in headers.values() for h in hs]:
        forms.setdefault(_naming_key(name), set()).add(name)
    clashes = [sorted(f) for f in forms.values() if len(f) > 1]
    if clashes:
        findings.append(Finding("deterministic", "naming", "major", "sheet and column names",
                                "Same concept, different names: " + "; ".join(" vs ".join(c) for c in clashes),
                                "Pick one name per concept and use it everywhere"))
    return findings


def check_features(spec: dict, build_dir: Path) -> list[Finding]:
    """Every feature the spec promises has visible evidence in the build."""
    findings = []
    workbooks: dict[str, object] = {}
    for feature in spec.get("features", []):
        evidence = feature["evidence"]
        path = build_dir / evidence["file"]
        problem = None
        if not path.is_file():
            problem = f"{evidence['file']} is missing"
        elif "sheet" in evidence:
            if evidence["file"] not in workbooks:
                workbooks[evidence["file"]] = load_workbook(path)
            wb = workbooks[evidence["file"]]
            ws = wb[evidence["sheet"]] if evidence["sheet"] in wb.sheetnames else None
            if ws is None:
                problem = f"sheet '{evidence['sheet']}' is missing"
            elif "column" in evidence and evidence["column"] not in {str(c.value).strip() for c in ws[1] if c.value}:
                problem = f"column '{evidence['column']}' is missing on {ws.title}"
            elif "label" in evidence and not any(isinstance(c.value, str) and c.value.strip() == evidence["label"]
                                                 for row in ws.iter_rows() for c in row):
                problem = f"'{evidence['label']}' is not shown on {ws.title}"
            elif len(ws._charts) < evidence.get("charts_min", 0):
                problem = f"{ws.title} has {len(ws._charts)} chart(s), needs {evidence['charts_min']}"
        if problem:
            findings.append(Finding("deterministic", "spec_features", "critical", evidence["file"],
                                    f"Promised feature '{feature['id']}' ({feature['description']}): {problem}",
                                    "Build the feature or stop promising it"))
    return findings


def check_docs_mention_files(spec: dict, doc_texts: dict[str, str]) -> list[Finding]:
    """The README and the PDF each name every other file in the download."""
    findings = []
    names = [f["name"] for f in spec["files"]]
    for doc, text in doc_texts.items():
        flat = re.sub(r"\s+", "", text)  # PDF extraction can break long names across lines
        missing = [n for n in names if n != doc and n not in text and n not in flat]
        if missing:
            findings.append(Finding("deterministic", "docs_mention_files", "major", doc,
                                    f"{doc} does not mention {', '.join(missing)}",
                                    "List every file in the download, by its exact name"))
    return findings


# Columns an example may leave empty: free-text notes nobody needs to see filled in.
_OPTIONAL_EXAMPLE_COLUMN = re.compile(r"\bnotes?\b|\bcomments?\b", re.I)


def check_example_fill(spec: dict) -> list[Finding]:
    """A filled-in example fills every entry column of each table it uses. A blank Vendor or Vehicle
    column reads as a broken product, in the workbook and in the listing pictures made from it."""
    findings = []
    files = {f["name"]: f for f in spec["files"]}
    for example in (f for f in spec["files"] if f.get("example_of")):
        base = files.get(example["example_of"], {})
        filled: dict[str, set[str]] = {}
        for ref in example.get("example_data", []):
            if "column" in ref and ref.get("value") not in (None, ""):
                filled.setdefault(ref["sheet"], set()).add(ref["column"])
        for sheet in base.get("sheets", []):
            if sheet.get("kind") != "table" or sheet["name"] not in filled:
                continue
            columns = sheet.get("column_defs") or sheet.get("columns") or []
            empty = [c["name"] for c in columns if not c.get("formula") and c["name"] not in filled[sheet["name"]]
                     and not _OPTIONAL_EXAMPLE_COLUMN.search(c["name"])]
            if empty:
                findings.append(Finding(
                    "deterministic", "example_fill", "major", f"{example['name']}:{sheet['name']}",
                    f"The example leaves {', '.join(empty)} empty on every row of {sheet['name']}",
                    f"Give the example rows realistic made-up values for {', '.join(empty)}"))
    return findings


def deterministic_checks(spec: dict, build_dir: Path, extra_words: set[str] | None = None) -> list[Finding]:
    extra_words = allowlist() if extra_words is None else extra_words
    texts: dict[str, str] = {}
    sheet_names: list[str] = []
    headers: dict[str, list[str]] = {}
    doc_texts: dict[str, str] = {}
    findings: list[Finding] = []
    for file_spec in spec["files"]:
        path = build_dir / file_spec["name"]
        if not path.is_file():
            continue  # Testing already failed missing files; features below will say so too
        try:
            if file_spec["kind"] == "xlsx":
                for sheet, entry in workbook_text(file_spec, path).items():
                    sheet_names.append(sheet)
                    if entry["headers"]:
                        headers[sheet] = entry["headers"]
                    texts[f"{file_spec['name']}:{sheet}"] = " ".join(
                        [sheet] + entry["headers"] + entry["text"] + entry["charts"])
                findings += [Finding("deterministic", p.check, "major", p.where, p.issue, p.suggestion)
                             for p in chart_checks(file_spec, path)]
            elif file_spec["kind"] == "pdf":
                doc_texts[file_spec["name"]] = texts[file_spec["name"]] = pdf_text(path)
            else:
                doc_texts[file_spec["name"]] = texts[file_spec["name"]] = path.read_text(encoding="utf-8")
        except Exception as error:  # unreadable for a buyer is a QA problem too
            findings.append(Finding("deterministic", "readable", "critical", file_spec["name"],
                                    f"Could not read {file_spec['name']}: {type(error).__name__}: {error}"))
    findings += check_spelling(texts, extra_words)
    findings += check_naming(sheet_names, headers)
    findings += check_features(spec, build_dir)
    findings += check_docs_mention_files(spec, doc_texts)
    findings += check_example_fill(spec)
    return findings


# --------------------------------------------------------------------------- scoring

def score_report(findings: list[Finding], *, pass_score: int, model_cap: int, vision_cap: int = 0) -> dict:
    """Deterministic findings deduct in full and can block; each AI pass deducts at most its cap."""
    deterministic = sum(SEVERITY_POINTS[f.severity] for f in findings if f.source not in AI_SOURCES)
    model = min(model_cap, sum(SEVERITY_POINTS[f.severity] for f in findings if f.source == "model"))
    vision = min(vision_cap, sum(SEVERITY_POINTS[f.severity] for f in findings if f.source == "vision"))
    score = max(0, 100 - deterministic - model - vision)
    blocking = [f for f in findings if f.source not in AI_SOURCES and f.severity in BLOCKING]
    decision = "PASS" if score >= pass_score and not blocking else "REWORK"
    return {"score": score, "decision": decision, "pass_score": pass_score, "deterministic_deduction": deterministic,
            "model_deduction": model, "model_deduction_cap": model_cap, "vision_deduction": vision,
            "vision_deduction_cap": vision_cap}


def required_changes(findings: list[Finding]) -> list[dict]:
    """REWORK brief items: every critical or major finding, deterministic first."""
    changes = []
    ordered = sorted(findings, key=lambda f: (f.source in AI_SOURCES, f.check))
    for index, finding in enumerate((f for f in ordered if f.severity in BLOCKING), start=1):
        changes.append({"id": f"qa-{finding.check.replace('_', '-')}-{index}",
                        "change": finding.suggestion or f"Fix: {finding.issue}",
                        "reason": f"[{finding.severity}] {finding.where}: {finding.issue}",
                        "acceptance": f"QA no longer reports this {finding.check} finding"})
    return changes


# --------------------------------------------------------------------------- department

Renderer = Callable[[list[Path], Path], RenderResult]


class QA:
    provider = "qa"

    def __init__(self, settings: Settings, client: httpx.Client | None = None, renderer: Renderer | None = None):
        self.settings = settings
        self.client = client
        self.renderer = renderer or (lambda workbooks, out_dir: render_workbooks(
            workbooks, out_dir, preference=settings.render_engine, soffice_path=settings.soffice_path))

    def __call__(self, product: Product, job: Job) -> StageResult:
        session = object_session(product)
        spec = product_spec(session, product, self.settings.products_dir)
        if spec is None:
            return StageResult(Outcome.FAIL, S.REWORK, f"QA: no specification for '{product.slug}'")
        build = session.scalar(select(Build).where(Build.product_id == product.id)
                               .order_by(Build.version.desc()).limit(1))
        if build is None:
            return StageResult(Outcome.FAIL, S.REWORK, "QA: no build recorded for this product")
        build_dir = Path(build.path)

        findings = deterministic_checks(spec, build_dir)
        provider = OllamaProvider(self.settings, session, client=self.client)
        completion = provider.generate(job, Department.QA, PROMPT.format(digest=product_digest(spec, build_dir)),
                                       system=SYSTEM, response_model=ModelReview)  # Ollama down: raises, job retries
        model_review: ModelReview = completion.data
        findings += [Finding("model", "clarity", f.severity, f.where, f.issue, f.suggestion)
                     for f in model_review.findings]

        render = self.renderer([build_dir / f["name"] for f in spec["files"] if f["kind"] == "xlsx"],
                               build_dir / RENDERS_DIR)
        findings += [Finding("render", "render", "info", "visual render", warning) for warning in render.warnings]
        findings += self._vision_pass(provider, job, render, spec)

        report = score_report(findings, pass_score=self.settings.qa_pass_score,
                              model_cap=self.settings.qa_model_max_deduction,
                              vision_cap=self.settings.qa_vision_max_deduction)
        report |= {"build_version": build.version, "clarity_score": model_review.clarity_score,
                   "model": completion.model, "render": render.as_dict(), "findings": [asdict(f) for f in findings]}
        counts = {s: sum(f.severity == s for f in findings) for s in ("critical", "major", "minor")}
        visual = (f"{len(render.charts)} chart(s) rendered by {render.engine} and reviewed" if render.engine
                  else "visual render skipped")
        summary = (f"QA {report['decision']} v{build.version}: score {report['score']}/100 "
                   f"(pass {report['pass_score']}); {counts['critical']} critical, {counts['major']} major, "
                   f"{counts['minor']} minor; model clarity {model_review.clarity_score}/100; {visual}")
        review = Review(product_id=product.id, job_id=job.id, build_id=build.id, department=Department.QA,
                        reviewer=f"qa+{provider.name}", model=completion.model, decision=report["decision"],
                        score=report["score"], summary=summary, findings_json=json.dumps(report["findings"]),
                        report_json=json.dumps(report))
        session.add(review)
        session.flush()
        log_event(session, "qa.report", actor="qa", product=product, job=job, message=summary,
                  data={"review_id": review.id, "score": report["score"], "decision": report["decision"]})
        data = {"review_id": review.id, "score": report["score"], "build_id": build.id, "version": build.version}
        if report["decision"] == "PASS":
            return StageResult(Outcome.PASS, S.RED_TEAM, summary, data)

        blocking = [f for f in findings if f.severity in BLOCKING]
        changes = required_changes(findings) or [{
            "id": "qa-score", "change": "Raise the QA score: address the minor findings in the QA report",
            "reason": f"Score {report['score']} is below the pass mark {report['pass_score']}", "acceptance": ""}]
        brief = create_brief(session, product, source="QA", actor="qa", summary=summary, required_changes=changes,
                             findings=[asdict(f) for f in blocking] or report["findings"], review=review, job=job,
                             build=build)
        top = "; ".join(f"[{f.severity}] {f.issue}" for f in blocking[:3])
        return StageResult(Outcome.FAIL, S.REWORK, f"{summary}. {top}".strip(), data | {"rework_brief_id": brief.id})

    def _vision_pass(self, provider: OllamaProvider, job: Job, render: RenderResult, spec: dict) -> list[Finding]:
        """One local-model call per chart image. Ollama down raises, and the job retries like the text pass."""
        findings = []
        roles = {f["name"]: "example" if f.get("example_of") else f.get("role", "template") for f in spec["files"]}
        for chart in render.charts:
            role = roles.get(chart.workbook, "template")
            state = VISION_STATES.get(role, VISION_STATES["template"])
            prompt = VISION_PROMPT.format(title=chart.title, sheet=chart.sheet, state=state)
            completion = provider.generate(job, Department.QA, prompt, system=VISION_SYSTEM,
                                           response_model=ChartObservation, images=[Path(chart.png).read_bytes()])
            where = f"{chart.sheet}: chart {chart.index} '{chart.title}' ({Path(chart.png).name})"
            findings += [Finding("vision", "chart_visual", severity, where, f"[{problem}] {issue}", suggestion)
                         for severity, problem, issue, suggestion in judge_observation(
                             completion.data, template=role != "example")]
        return findings
